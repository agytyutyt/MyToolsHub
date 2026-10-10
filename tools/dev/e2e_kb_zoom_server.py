# -*- coding: utf-8 -*-
"""E2E 临时服务器：验证知识库预览缩放（xls 卡片同步放大 + 百分比输入）。
隔离数据根 JZTOOLS_DATA_ROOT，假 admin 会话，托管插件前端，预置一份 xls。
端口 5177。用法：python e2e_kb_zoom_server.py serve|seed-only
"""
import importlib
import io
import os
import shutil
import subprocess
import sys
import tempfile
import types

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))  # tools/dev/ → 仓库根
DATA = os.path.join(REPO, ".workbuddy", "tmp", "e2e-kb-data")   # 隔离数据根，已 gitignore
os.environ["JZTOOLS_DATA_ROOT"] = DATA
os.makedirs(DATA, exist_ok=True)
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
os.environ["XHR_SOFFICE"] = SOFFICE

sys.path.insert(0, REPO)

# ---- 伪造 jztools_admin（routes 顶层 try-import）----
# E2E_KB_USER=b 时切换为「另一单位管理员 kb-b」（分类可见范围 E2E 用：
# 非 super、U2/D9；对 tester 建的 unit/dept/private 分类均不可见）；
# 默认仍是超管 tester（缩放验证等既有用法不受影响）。
if os.environ.get("E2E_KB_USER") == "b":
    _SESSION = {"username": "kb-b", "name": "乙管理员", "super_admin": False,
                "role_id": "role-admin", "unit_id": "U2", "department_id": "D9"}
else:
    _SESSION = {"username": "tester", "name": "测试管理员", "super_admin": True,
                "role_id": "role-admin", "unit_id": "U1", "department_id": "D1"}
admin = types.ModuleType("jztools_admin")
admin_routes = types.ModuleType("jztools_admin.routes")
admin_routes.get_session_user = lambda: dict(_SESSION)
admin_routes.set_operation = lambda op: None
admin.routes = admin_routes
sys.modules["jztools_admin"] = admin
sys.modules["jztools_admin.routes"] = admin_routes

# routes 从框架门面取身份（jz_api.get_session_user）——必须按真实注册路径补 provider，
# 否则一切需登录接口 401（口径同 test_routes_preview.py 的修法）。
import jz_api  # noqa: E402

jz_api.register_providers(get_session_user=admin_routes.get_session_user,
                          set_operation=admin_routes.set_operation)

pkg = types.ModuleType("kb_backend_e2e")
pkg.__path__ = [os.path.join(REPO, "plugins", "knowledge-base", "backend")]
sys.modules["kb_backend_e2e"] = pkg
routes = importlib.import_module("kb_backend_e2e.routes")

from flask import Flask, send_from_directory  # noqa: E402

FRONTEND = os.path.join(REPO, "plugins", "knowledge-base", "frontend")

app = Flask(__name__)
routes.register(app)


@app.route("/")
def home():
    return send_from_directory(FRONTEND, "index.html")


@app.route("/plugin/knowledge-base/<path:filename>")
def assets(filename):
    return send_from_directory(FRONTEND, filename)


def make_xls_bytes():
    """openpyxl 造 xlsx → soffice 转 xls（唯一 profile，避免单实例冲突）。"""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "数据"
    ws["A1"] = "区域"; ws["B1"] = "销售额"
    ws["A2"] = "华东"; ws["B2"] = 1234567.891
    for i in range(3, 40):
        ws.cell(row=i, column=1, value="明细行%d" % i)
        ws.cell(row=i, column=2, value=i * 100)
    buf = io.BytesIO()
    wb.save(buf)
    work = tempfile.mkdtemp(prefix="kbe2e-")
    try:
        sp = os.path.join(work, "book.xlsx")
        with open(sp, "wb") as f:
            f.write(buf.getvalue())
        prof = os.path.join(work, "profile")
        subprocess.run(
            [SOFFICE, "--headless",
             "-env:UserInstallation=file:///" + prof.replace("\\", "/"),
             "--convert-to", "xls", "--outdir", work, sp],
            capture_output=True, timeout=120)
        dp = os.path.join(work, "book.xls")
        if not os.path.isfile(dp):
            raise RuntimeError("soffice 转换 xls 失败: " +
                               subprocess.list2cmdline([SOFFICE]))
        with open(dp, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)


def seed():
    xls_bytes = make_xls_bytes()
    c = app.test_client()
    r = c.post("/api/knowledge-base/files", data={
        "file": (io.BytesIO(xls_bytes), "季度报表.xls"), "name": "季度报表"},
        content_type="multipart/form-data")
    print("seed upload:", r.status_code, r.get_json())
    r = c.get("/api/knowledge-base/files")
    items = (r.get_json() or {}).get("files") or []
    print("seed files:", [(x.get("name"), x.get("ext")) for x in items])
    # 触发一次预览渲染，确认后端 OK
    if items:
        r = c.get("/api/knowledge-base/files/%s/preview" % items[0]["id"])
        d = r.get_json() or {}
        print("seed preview:", r.status_code, "kind=", d.get("kind"),
              "html=", len(d.get("html") or ""))


def serve():
    print("serving on http://127.0.0.1:5177  data_root=", DATA, flush=True)
    app.run(host="127.0.0.1", port=5177, threaded=True)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "seed-only":
        seed()
    else:
        serve()
