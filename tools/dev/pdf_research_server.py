# -*- coding: utf-8 -*-
"""E2E 临时服务器：排查/回归 PDF 预览清晰度（画布位图 vs 显示设备像素）。

背景：用户反馈"PDF 预览分辨率低、发糊"。实测 dpr 一直是对的，病根是**画布没落在设备
像素网格上**（居中 + A4 小数高度算出小数相位 → 合成器对整张画布双线性重采样）。
判定口径：位图尺寸 == 显示设备像素、两轴相位（`rect.left*dpr % 1` 与 top 同）≈ 0。

用法：
    python pdf_research_server.py seed-only     # 造一份多页 A4 测试 PDF 并入库
    python pdf_research_server.py serve         # 起服务（端口 5178，隔离数据根）

配套探测（`.workbuddy/pdf-research/` 下留产出）：
  1) 探针页 http://127.0.0.1:5178/plugin/knowledge-base/index-probe.html?noio[&id=<文件id>]
     直接调 KBReader.render 打开 PDF，把 dpr / 位图尺寸 / 显示设备像素 / 两轴相位
     POST 回 /probe-report（收在 pdf-research/probe-report.jsonl）；
     `?noio` 置空 IntersectionObserver 走"顺序渲染全部页"分支（headless 下 IO 不回调）。
  2) 截图出图（load 事件被 /slow.png 拖住 12s，保证截图落在 PDF 渲染完成之后）：
     chrome --headless=new --force-device-scale-factor=1.25 --window-size=1280,900 \
       --screenshot=shot.png URL?noio
  3) 客观锐度：shot.png 与画布原始位图（rawcrop 里的 dataUrl）用 PIL/numpy 算
     Laplacian 边缘能量对比；同位置前后截图对比可直接得出"锐度提升 N%"。
"""
import importlib
import io
import os
import sys
import types

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
DATA = os.path.join(REPO, ".workbuddy", "e2e-pdf-data")
os.environ["JZTOOLS_DATA_ROOT"] = DATA
os.makedirs(DATA, exist_ok=True)

SOFFICE = os.environ.get("XHR_SOFFICE", r"C:\Program Files\LibreOffice\program\soffice.exe")

sys.path.insert(0, REPO)

admin = types.ModuleType("jztools_admin")
admin_routes = types.ModuleType("jztools_admin.routes")
admin_routes.get_session_user = lambda: {
    "username": "tester", "name": "测试管理员", "super_admin": True,
    "role_id": "role-admin", "unit_id": "", "department_id": ""}
admin_routes.set_operation = lambda op: None
admin.routes = admin_routes
sys.modules["jztools_admin"] = admin
sys.modules["jztools_admin.routes"] = admin_routes

pkg = types.ModuleType("kb_backend_pdf")
pkg.__path__ = [os.path.join(REPO, "plugins", "knowledge-base", "backend")]
sys.modules["kb_backend_pdf"] = pkg
routes = importlib.import_module("kb_backend_pdf.routes")

from flask import Flask, send_from_directory, request  # noqa: E402

FRONTEND = os.path.join(REPO, "plugins", "knowledge-base", "frontend")

app = Flask(__name__)
routes.register(app)


@app.route("/")
def home():
    return send_from_directory(FRONTEND, "index.html")


@app.route("/plugin/knowledge-base/index-probe.html")
def probe_page():
    """index.html + 探针脚本：直接调 KBReader.render 打开预置 PDF，
    测量值 POST 回 /probe-report（headless 下 --dump-dom 与 virtual-time 会卡 worker）。"""
    with open(os.path.join(FRONTEND, "index.html"), encoding="utf-8") as f:
        html = f.read()
    probe = """
<script>
window.__noIO = /noio/.test(location.search);
function report(payload) {
  try {
    fetch("/probe-report", { method: "POST", headers: { "Content-Type": "application/json" },
      body: JSON.stringify(payload) });
  } catch (e) { /* ignore */ }
}
(function () {
  function out(lines) {
    var pre = document.getElementById("probe-out");
    if (!pre) {
      pre = document.createElement("pre");
      pre.id = "probe-out";
      pre.style.cssText = "position:fixed;left:0;bottom:0;z-index:9999;font:11px monospace;background:#fff";
      document.body.appendChild(pre);
    }
    pre.textContent = lines.join("\\n");
  }
    // ---- 合成器重采样检测：参考图 + 画布原始位图回传 ----
    window.__probe4 = function () {
      var c = document.querySelector(".pdf-page-box canvas");
      if (!c) { report({ phase: "rawcrop", err: "no canvas" }); return; }
      var dpr = window.devicePixelRatio;
      var r = c.getBoundingClientRect();
      // 位图坐标区域（画布内文字区）→ 对应屏幕设备像素位置 = rect*dpr + 位图偏移
      var cx = 40, cy = 60, cw = 400, ch = 160;
      var tmp = document.createElement("canvas");
      tmp.width = cw; tmp.height = ch;
      tmp.getContext("2d").drawImage(c, cx, cy, cw, ch, 0, 0, cw, ch);
      // 参考：1 设备像素宽的竖条纹画布（CSS 尺寸 = 位图/dpr），合成器若 1:1 合成则条纹不褪色
      var ref = document.createElement("canvas");
      ref.width = 240; ref.height = 60;
      ref.style.cssText = "position:fixed;left:0;top:0;z-index:99999;width:"
        + (240 / dpr) + "px;height:" + (60 / dpr) + "px";
      var rc = ref.getContext("2d");
      for (var x = 0; x < 240; x++) {
        rc.fillStyle = (x % 2) ? "#000" : "#fff";
        rc.fillRect(x, 0, 1, 60);
      }
      document.body.appendChild(ref);
      // 受控 A/B：相同内容，A 落在整数设备像素、B 落在半像素（检验亚像素定位是否致糊）
      var dprT = dpr;
      function mkStrip(leftCss, topCss) {
        var cv = document.createElement("canvas");
        cv.width = 300; cv.height = 60;
        cv.style.cssText = "position:fixed;left:" + leftCss + "px;top:" + topCss
          + "px;z-index:99999;width:" + (300 / dprT) + "px;height:" + (60 / dprT) + "px";
        var g = cv.getContext("2d");
        g.fillStyle = "#fff"; g.fillRect(0, 0, 300, 60);
        g.fillStyle = "#000"; g.font = "20px Arial";
        g.fillText("Hamburgefonstiv 1234567890", 2, 24);
        g.font = "14px Arial";
        g.fillText("fine print sharpness probe 8pt", 2, 46);
        document.body.appendChild(cv);
        return cv;
      }
      // 对齐调试：对每个 PDF 画布做"读相位 → 校正 → 复测"，逐步上报
      window.__snapDebug = function () {
        var cs = document.querySelectorAll(".pdf-page-box canvas");
        var out2 = [];
        for (var i = 0; i < cs.length; i++) {
          (function (c, idx) {
            var d = window.devicePixelRatio, steps = [];
            for (var k = 0; k < 4; k++) {
              var r = c.getBoundingClientRect();
              var fx = (r.left * d) % 1;
              steps.push({ step: k, left: Math.round(r.left * 100000) / 100000,
                           phase: Math.round(fx * 100000) / 100000,
                           margin: c.style.marginLeft || "(空)" });
              if (fx < 0.002 || fx > 0.998) break;
              c.style.marginLeft = ((parseFloat(c.style.marginLeft) || 0) - fx / d) + "px";
            }
            out2.push({ page: idx + 1, steps: steps });
          })(cs[i], i);
        }
        report({ phase: "snapdebug", dpr: window.devicePixelRatio, pages: out2 });
      };
      mkStrip(0, 72);                      // 设备 x = 0（整数）
      mkStrip(100.4, 72);                  // 设备 x = 125.5（半像素）
      report({
        phase: "rawcrop",
        dpr: dpr,
        rect: [Math.round(r.left * 100) / 100, Math.round(r.top * 100) / 100,
               Math.round(r.width * 100) / 100, Math.round(r.height * 100) / 100],
        bitmap: [c.width, c.height],
        crop: [cx, cy, cw, ch],
        scrollY: window.pageYOffset,
        dataUrl: tmp.toDataURL("image/png"),
        refRect: [0, 0, 240 / dpr, 60 / dpr],
        refBitmap: [240, 60]
      });
    };
    window.__probeDone = false;
    window.__probe = function () {
    var res = {
      dpr: window.devicePixelRatio,
      vvScale: window.visualViewport ? window.visualViewport.scale : null,
      innerW: window.innerWidth,
      canvases: []
    };
    var cs = document.querySelectorAll(".pdf-page-box canvas");
    for (var i = 0; i < cs.length; i++) {
      var c = cs[i], r = c.getBoundingClientRect();
      res.canvases.push({
        bitmap: [c.width, c.height],
        css: [Math.round(r.width * 100) / 100, Math.round(r.height * 100) / 100],
        bmpPerCssPx: Math.round((c.width / r.width) * 10000) / 10000,
        inline: [c.style.width, c.style.height]
      });
    }
    var box = document.querySelector(".pdf-page-box");
    if (box) {
      var br = box.getBoundingClientRect();
      res.box = [Math.round(br.width * 100) / 100, Math.round(br.height * 100) / 100];
    }
    var out2 = ["PROBE " + JSON.stringify(res)];
    // 画布位图内的文字行高（判断实际光栅分辨率）
    if (cs.length) {
      var ctx = cs[0].getContext("2d");
      out2.push("BITMAP_W " + cs[0].width + " BITMAP_H " + cs[0].height);
    }
    out(out2);
    report(res);
    return out2[0];
  };
  window.__probe2 = function (tag) {
    var cs = document.querySelectorAll(".pdf-page-box canvas");
    var res = {
      phase: tag,
      dpr: window.devicePixelRatio,
      zoomInput: (document.getElementById("zoom-input") || {}).value,
      canvases: []
    };
    for (var i = 0; i < cs.length; i++) {
      var c = cs[i], r = c.getBoundingClientRect();
      res.canvases.push({
        bitmap: [c.width, c.height],
        css: [Math.round(r.width * 100) / 100, Math.round(r.height * 100) / 100],
        bmpPerCssPx: Math.round((c.width / r.width) * 10000) / 10000
      });
    }
    var box = document.querySelector(".pdf-page-box");
    if (box) {
      var br = box.getBoundingClientRect();
      res.box = [Math.round(br.width * 100) / 100, Math.round(br.height * 100) / 100];
    }
    report(res);
    return res;
  };
  function boot() {
    if (!window.KBReader) { setTimeout(boot, 200); return; }
    document.getElementById("view-list").className = "view hidden";
    document.getElementById("view-reader").className = "view";
    document.getElementById("reader-status").className = "reader-status hidden";
    // headless 下 IntersectionObserver 回调不触发：置空走"顺序渲染全部页"分支，
    // 与 IO 分支共用同一个 renderPage（位图尺寸与 scale/dpr 的计算完全一致）
    if (window.__noIO) window.IntersectionObserver = undefined;
    var file = { id: (new URLSearchParams(location.search)).get("id") || "f-0d3b82c2",
                 ext: "pdf", name: "probe", original_name: "sharp.pdf" };
    window.KBReader.render(file, {
      onReady: function () { window.__probe(); },
      onPage: function () { window.__probe(); },
      onFail: function (m) { out(["PROBE_FAIL " + m]); }
    });
    // 轮询直到画布出现（headless 下 IO 触发时机不定）
    var tries = 0;
    function measure(tag, then) {
      var cs = document.querySelectorAll(".pdf-page-box canvas");
      if (!cs.length) return false;
      var r = window.__probe2(tag);
      if (then) then();
      return r;
    }
    (function poll() {
      tries++;
      if (measure("initial", phase2)) return;
      if (tries > 40) { window.__probe(); return; }
      setTimeout(poll, 500);
    })();
    var probe4Done = false;
    (function poll4() {
      var cs = document.querySelectorAll(".pdf-page-box canvas");
      if (cs.length && !probe4Done) { probe4Done = true; setTimeout(function () { window.__probe4(); }, 300); setTimeout(function () { window.__snapDebug(); }, 5000); return; }
      setTimeout(poll4, 500);
    })();
    // 持续上报最新布局（截图在 load 后触发，中间可能因状态条隐藏/滚动条出现而移位）
    setInterval(function () {
      var cs = document.querySelectorAll(".pdf-page-box canvas");
      if (!cs.length) return;
      var dprN = window.devicePixelRatio;
      var pages = [];
      for (var i = 0; i < cs.length; i++) {
        var r = cs[i].getBoundingClientRect();
        pages.push({ page: i + 1,
                     leftPhase: Math.round(((r.left * dprN) % 1) * 10000) / 10000,
                     topPhase: Math.round(((r.top * dprN) % 1) * 10000) / 10000,
                     bitmap: [cs[i].width, cs[i].height],
                     devSize: [Math.round(r.width * dprN * 1000) / 1000,
                               Math.round(r.height * dprN * 1000) / 1000] });
      }
      report({ phase: "rect", pages: pages, scrollY: window.pageYOffset, dpr: dprN });
    }, 400);
    var phase2Done = false;
    function phase2() {
      if (phase2Done) return;
      phase2Done = true;
      // 缩放路径：setZoom(2.0) 后应重渲染，位图/CSS 比保持 = dpr
      try { window.KBReader.setZoom(2.0); } catch (e) { report({ phase: "zooomErr", err: String(e) }); return; }
      var t2 = 0;
      (function poll2() {
        t2++;
        var c = document.querySelector(".pdf-page-box canvas");
        if (c) {
          var r = c.getBoundingClientRect();
          report({ phase: "afterZoom2",
                   topPhase: (r.top * window.devicePixelRatio) % 1,
                   leftPhase: (r.left * window.devicePixelRatio) % 1,
                   scrollY: window.pageYOffset,
                   scrollPhase: (window.pageYOffset * window.devicePixelRatio) % 1,
                   canvases: [{ bitmap: [c.width, c.height],
                                css: [Math.round(r.width), Math.round(r.height)] }] });
          return;
        }
        if (t2 > 40) { report({ phase: "afterZoom2", canvases: [] }); return; }
        setTimeout(poll2, 500);
      })();
    }
  }
  // 立即启动（不要等 load：load 被 /slow.png 拖住，用于让 --screenshot 晚于 PDF 渲染）
  boot();
})();
</script>
"""
    html = html.replace("</body>", probe + "</body>", 1)
    html = html.replace("<body>", '<body><img src="/slow.png" width="1" height="1" alt="" '
                                  'style="position:fixed;left:-10px;top:-10px">', 1)
    return html


@app.route("/plugin/knowledge-base/index-debug.html")
def debug_page():
    """同一份 index.html，仅在 <head> 后注入错误采集器（不改插件源码）。
    必须挂在 /plugin/knowledge-base/ 下，页面里的 ./reader.js 相对路径才解析正确。"""
    with open(os.path.join(FRONTEND, "index.html"), encoding="utf-8") as f:
        html = f.read()
    hook = (
        "<script>window.__errs=[];"
        "window.addEventListener('error',function(e){window.__errs.push("
        "'ERR '+e.message+' @'+(e.filename||'')+':'+e.lineno)});"
        "window.addEventListener('unhandledrejection',function(e){window.__errs.push("
        "'REJ '+String((e.reason&&e.reason.message)||e.reason))});</script>"
    )
    html = html.replace("<head>", "<head>" + hook, 1)
    return html


@app.route("/plugin/knowledge-base/<path:filename>")
def assets(filename):
    return send_from_directory(FRONTEND, filename)


REPORT_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "pdf-research", "probe-report.jsonl")


@app.route("/slow.png")
def slow_png():
    """拖住 load 事件：让 headless 的 --screenshot（load 后触发）落在 PDF 渲染完成之后。"""
    import time as _t
    from flask import Response
    _t.sleep(12)
    png = bytes.fromhex(
        "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c4"
        "890000000a49444154789c6300010000050001"
        "0d0a2db40000000049454e44ae426082")
    return Response(png, mimetype="image/png")


@app.route("/probe-report", methods=["POST"])
def probe_report():
    import json as _json
    rec = request.get_json(silent=True) or {}
    rec["_ua"] = request.headers.get("User-Agent", "")[:80]
    with open(REPORT_PATH, "a", encoding="utf-8") as f:
        f.write(_json.dumps(rec, ensure_ascii=False) + "\n")
    return {"ok": True}


def make_a4_pdf_bytes(pages=3):
    """python-docx 造多页 A4（含 8pt 小字，便于肉眼/像素级判清晰度）→ soffice 转 PDF。
    A4 是关键：841.89pt × 1.25 = 1052.36px 的小数高度正是历史上线相位失配的来源。"""
    import subprocess
    import tempfile
    import shutil
    import docx
    from docx.shared import Mm, Pt
    from docx.enum.text import WD_BREAK

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    for pg in range(pages):
        if pg:
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        p = d.add_paragraph()
        p.add_run("Fine print 8pt: The quick brown fox jumps over the lazy dog. 0123456789"
                  ).font.size = Pt(8)
        for i in range(16):
            d.add_paragraph("Page %d line %d body text 10.5pt sharpness sample." % (pg + 1, i + 1))
    work = tempfile.mkdtemp(prefix="kbpdf-")
    try:
        src = os.path.join(work, "probe.docx")
        d.save(src)
        subprocess.run(
            [SOFFICE, "--headless",
             "-env:UserInstallation=file:///" + os.path.join(work, "profile").replace("\\", "/"),
             "--convert-to", "pdf", "--outdir", work, src],
            capture_output=True, timeout=180)
        out = os.path.join(work, "probe.pdf")
        if not os.path.isfile(out):
            raise RuntimeError("soffice 转换 PDF 失败：" + SOFFICE)
        with open(out, "rb") as f:
            return f.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)


def seed():
    data = make_a4_pdf_bytes()
    c = app.test_client()
    r = c.post("/api/knowledge-base/files", data={
        "file": (io.BytesIO(data), "probe-a4.pdf"), "name": "清晰度测试"},
        content_type="multipart/form-data")
    print("seed upload:", r.status_code, r.get_json())
    r = c.get("/api/knowledge-base/files")
    print("seed files:", [(x.get("id"), x.get("name"), x.get("ext"))
                          for x in ((r.get_json() or {}).get("files") or [])])


def serve():
    print("serving on http://127.0.0.1:5178  data_root=", DATA, flush=True)
    app.run(host="127.0.0.1", port=5178, threaded=True)


if __name__ == "__main__":
    if len(sys.argv) > 1 and sys.argv[1] == "seed-only":
        seed()
    else:
        serve()
