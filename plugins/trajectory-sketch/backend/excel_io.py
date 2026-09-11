# -*- coding: utf-8 -*-
"""表格读写：轨迹表读取 + 速写报告工作簿输出。

职责边界（与引擎严格分离）
--------------------------
- **读**：Excel/CSV → 二维表 → **JSON 安全化**（``datetime`` 转字符串、``NaN`` 转空串），
  因为这张表要经 ``/api/file-filter/apply`` 走一次 JSON 往返，不能带非序列化类型；
- **写**：引擎结果 dict → 5 个 sheet 的 ``.xlsx`` 速写报告。

引擎（``engine/``）完全不感知本模块，也不感知"Excel"这个概念——它只处理二维表与 dict。
"""
from __future__ import annotations

import csv
import os
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Sequence, Tuple

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except Exception:  # pragma: no cover
    openpyxl = None
    OPENPYXL_AVAILABLE = False

try:
    import xlrd
    XLRD_AVAILABLE = True
except Exception:  # pragma: no cover
    xlrd = None
    XLRD_AVAILABLE = False

ALLOWED_EXTS = {"xlsx", "xls", "csv"}
MAX_ROWS = 100000
MAX_COLS = 256
_BOM = "\ufeff"


class TableError(Exception):
    """表格读写异常，信息直接展示给前端。"""


# --------------------------------------------------------------------------
# 单元格 JSON 安全化
# --------------------------------------------------------------------------

def json_safe(v: Any) -> Any:
    """把单元格值转为 JSON 可序列化类型（时间转字符串、NaN 转空串）。"""
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        if isinstance(v, datetime):
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return "" if v != v else v            # NaN → 空串
    if isinstance(v, (int, str, bool)):
        return v
    return str(v)


def clean_text(v: Any) -> str:
    if v is None:
        return ""
    s = v if isinstance(v, str) else str(v)
    return s.replace(_BOM, "").replace("\u3000", " ").strip()


# --------------------------------------------------------------------------
# 读取
# --------------------------------------------------------------------------

def read_rows(path: str, filename: str) -> List[List[Any]]:
    """读取表格为二维表（首行表头），单元格已做 JSON 安全化。"""
    ext = (filename or "").lower().rsplit(".", 1)[-1]
    if ext not in ALLOWED_EXTS:
        raise TableError("不支持的文件格式：.%s（仅支持 xlsx / xls / csv）" % ext)
    if ext == "csv":
        raw = _read_csv(path)
    elif ext == "xlsx":
        raw = _read_xlsx(path)
    else:
        raw = _read_xls(path)

    if not raw:
        raise TableError("表格为空。")
    raw = [list(r) for r in raw]
    while raw and all(clean_text(v) == "" for v in raw[-1]):
        raw.pop()
    if len(raw) > MAX_ROWS + 1:
        raise TableError("表格行数超过上限（最多 %d 行数据）。" % MAX_ROWS)
    if len(raw[0]) > MAX_COLS:
        raise TableError("表格列数超过上限（最多 %d 列）。" % MAX_COLS)
    if not any(clean_text(h) for h in raw[0]):
        raise TableError("表格缺少表头：首行必须为字段名称。")
    return raw


def _read_xlsx(path: str) -> List[List[Any]]:
    if not OPENPYXL_AVAILABLE:
        raise TableError("后端缺少 openpyxl，无法解析 .xlsx（请执行：pip install openpyxl）")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise TableError("Excel 文件解析失败（%s）。" % type(exc).__name__)
    try:
        ws = wb.active
        return [[json_safe(v) for v in row] for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _read_xls(path: str) -> List[List[Any]]:
    if not XLRD_AVAILABLE:
        raise TableError("后端缺少 xlrd，无法解析 .xls（请执行：pip install xlrd）")
    try:
        book = xlrd.open_workbook(path)
    except Exception as exc:
        raise TableError("Excel 文件解析失败（%s）。" % type(exc).__name__)
    sheet = book.sheet_by_index(0)
    out: List[List[Any]] = []
    for r in range(sheet.nrows):
        row = []
        for c in range(sheet.ncols):
            cell = sheet.cell(r, c)
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    row.append(xlrd.xldate_as_datetime(cell.value, book.datemode)
                               .strftime("%Y-%m-%d %H:%M:%S"))
                except Exception:
                    row.append(cell.value)
            else:
                row.append(json_safe(cell.value))
        out.append(row)
    return out


def _read_csv(path: str) -> List[List[Any]]:
    for enc in ("utf-8-sig", "gbk"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                rows = [list(r) for r in csv.reader(f)]
            if rows:
                return rows
            raise TableError("CSV 文件为空。")
        except UnicodeDecodeError:
            continue
        except TableError:
            raise
        except Exception as exc:
            raise TableError("CSV 解析失败（%s）。" % type(exc).__name__)
    raise TableError("无法识别的文件编码（请另存为 UTF-8 或 GBK 编码的 CSV）。")


def preview_headers(rows: Sequence[Sequence[Any]], limit: int = 200) -> List[str]:
    """表头预览（前端"上传后看到的字段"）。"""
    if not rows:
        return []
    return [clean_text(h) for h in list(rows[0])[:limit]]


# --------------------------------------------------------------------------
# 速写报告工作簿
# --------------------------------------------------------------------------

SHEET_REPORT = "速写报告"
SHEET_SUMMARY = "速写摘要"
SHEET_STAYS = "停留点"
SHEET_TRIPS = "出行段"
SHEET_QUALITY = "数据质量"


def write_report_workbook(path: str, result: Dict[str, Any]) -> str:
    """把引擎结果写成 5 个 sheet 的速写报告工作簿，返回落盘路径。"""
    if not OPENPYXL_AVAILABLE:
        raise TableError("后端缺少 openpyxl，无法生成报告（请执行：pip install openpyxl）")
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb = openpyxl.Workbook()

    report = result.get("report") or {}
    text_by_user: Dict[str, str] = report.get("text_by_user") or {}
    summary: Dict[str, Any] = report.get("summary") or {}

    # ---- sheet 1：速写报告（正文，列宽 100，自动换行，可直接复制）----
    ws = wb.active
    ws.title = SHEET_REPORT
    ws.append(["USERNUM", "guiji_result"])
    for user, text in text_by_user.items():
        ws.append([user or "（未标注号码）", text])
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 100
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    _bold_header(ws)

    # ---- sheet 2：速写摘要 ----
    ws2 = wb.create_sheet(SHEET_SUMMARY)
    ws2.append(["指标", "数值"])
    for k, v in summary.items():
        ws2.append([_label(k), _cell(v)])
    for user, s in (report.get("user_summaries") or {}).items():
        ws2.append(["号码 %s" % (user or "（未标注号码）"), s])
    _autofit(ws2)

    # ---- sheet 3 / 4：停留点 / 出行段 ----
    _write_dicts(wb.create_sheet(SHEET_STAYS), result.get("stays") or [])
    _write_dicts(wb.create_sheet(SHEET_TRIPS), result.get("trips") or [])

    # ---- sheet 5：数据质量 ----
    ws5 = wb.create_sheet(SHEET_QUALITY)
    quality: Dict[str, Any] = result.get("quality") or {}
    ws5.append(["指标", "数值"])
    for k, v in quality.items():
        if isinstance(v, dict):
            for k2, v2 in v.items():
                ws5.append(["%s · %s" % (_label(k), k2), _cell(v2)])
        elif isinstance(v, list):
            ws5.append([_label(k), "；".join(_cell(x) for x in v)])
        else:
            ws5.append([_label(k), _cell(v)])
    # 过滤阶段摘要（谁保留了哪些字段）与引擎警告
    flt = result.get("filter") or {}
    if flt:
        ws5.append(["", ""])
        ws5.append(["过滤模式", "大模型辅助" if flt.get("mode") == "llm" else "硬过滤"])
        ws5.append(["过滤后保留字段", "；".join(
            "%s%s" % (i.get("column", ""), ("（匹配 %s）" % i.get("matched")) if i.get("matched") else "")
            for i in (flt.get("kept") or []))])
        ws5.append(["过滤后删除字段", "；".join(str(x) for x in (flt.get("removed") or []))])
        ws5.append(["后处理替换次数", _cell(flt.get("replace_count"))])
    for w in (result.get("warnings") or []):
        ws5.append(["提示", w])
    _autofit(ws5)

    wb.save(path)
    return path


def _write_dicts(ws, items: List[Dict[str, Any]]) -> None:
    if not items:
        ws.append(["（无）"])
        return
    cols: List[str] = []
    for it in items:
        for k in it.keys():
            if k not in cols:
                cols.append(k)
    ws.append(cols)
    for it in items:
        ws.append([_cell(it.get(c)) for c in cols])
    _bold_header(ws)
    _autofit(ws)


def _bold_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"


def _autofit(ws, max_width: int = 42) -> None:
    """按前 200 行内容估算列宽（1~max_width）。"""
    widths: Dict[int, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 200)):
        for cell in row:
            v = "" if cell.value is None else str(cell.value)
            width = max(len(v), 6)
            widths[cell.column] = max(widths.get(cell.column, 0), width)
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(width + 2, max_width)


def _cell(v: Any) -> Any:
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        return "；".join(_cell(x) for x in v)
    if isinstance(v, dict):
        return "；".join("%s=%s" % (k, _cell(x)) for k, x in v.items())
    return v


_LABELS = {"net_m": "净位移(米)", "cum_m": "累计位移(米)", "kind": "判定",
           "confidence": "置信度", "lon": "经度", "lat": "纬度"}


def _label(k: str) -> str:
    return _LABELS.get(str(k), str(k))
