# -*- coding: utf-8 -*-
"""信息传输插件端到端回归基线（TODO T01）。

对 9 类格式样例走「封装 → 二维码（静态/视频）→ 解码 → 还原」全链路：
- 原件传输（fmt=file）：导出字节 SHA-256 必须与源文件一致；
- 精简传输：还原的数据结构（文本/二维数组）必须与提取结果一致。

样例来源：优先复用 `.workbuddy/tmp/it_bench/` 下的实测样例（含 soffice 生成的
doc/xls/ppt/pdf）；缺失时本地生成 txt/md/csv/xlsx/docx，soffice 可用则补生成
doc/xls/ppt/pdf，仍缺则跳过对应用例（标记 skip，不算失败）。

运行：`python test_roundtrip.py`（隔离数据根 JZTOOLS_DATA_ROOT 指向临时目录，
不触碰真实数据 ~/.jztoolshub）。
"""
import base64
import hashlib
import importlib.util
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import zipfile
import zlib

REPO = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
sys.path.insert(0, REPO)

TMP_ROOT = tempfile.mkdtemp(prefix="it_roundtrip_")
os.environ["JZTOOLS_DATA_ROOT"] = TMP_ROOT  # 隔离真实数据根（routes 导入前设置）

# 动态加载插件路由模块
_spec = importlib.util.spec_from_file_location(
    "it_routes_test", os.path.join(REPO, "plugins", "info-transfer",
                                   "backend", "routes.py"))
it = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(it)

BENCH = os.path.join(REPO, ".workbuddy", "tmp", "it_bench")
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"

SAMPLES = ["sample.txt", "sample.md", "sample.csv", "sample.xlsx",
           "sample.docx", "sample.doc", "sample.xls", "sample.ppt", "sample.pdf"]
SKIP = []          # 找不到样例的格式
RESULTS = []


# ---------------- 样例准备 ----------------

def _make_basic_samples():
    """用库生成 txt/md/csv/xlsx/docx 样例（bench 样例缺失时的兜底）。"""
    import openpyxl
    import docx as docx_mod
    text = ("# 回归测试样例\n\n" + "".join(
        f"第 {i} 段：本项目旨在解决内网环境下跨网闸的数据交换问题，通过二维码单向通道实现小体积文档的可靠传输。\n"
        for i in range(1, 41)))
    rows = [["序号", "部门", "姓名", "工号"]] + [
        [r, ["综合管理部", "技术研发部", "市场运营部"][r % 3],
         ["张伟", "王芳", "李娜"][r % 3], f"GH{20260000 + r}"] for r in range(1, 101)]
    p = lambda n: os.path.join(BENCH, n)
    os.makedirs(BENCH, exist_ok=True)
    if not os.path.exists(p("sample.txt")):
        open(p("sample.txt"), "w", encoding="utf-8").write(text)
    if not os.path.exists(p("sample.md")):
        open(p("sample.md"), "w", encoding="utf-8").write(text)
    if not os.path.exists(p("sample.csv")):
        open(p("sample.csv"), "w", encoding="utf-8").write(
            "\n".join(",".join(str(v) for v in row) for row in rows))
    if not os.path.exists(p("sample.xlsx")):
        wb = openpyxl.Workbook(); ws = wb.active
        for row in rows:
            ws.append(row)
        wb.save(p("sample.xlsx"))
    if not os.path.exists(p("sample.docx")):
        d = docx_mod.Document()
        for i in range(1, 21):
            d.add_heading(f"第 {i} 节", 2)
            d.add_paragraph("本项目旨在解决内网环境下跨网闸的数据交换问题，通过二维码单向通道实现可靠传输。")
        d.save(p("sample.docx"))


def _soffice_convert(src, target):
    out = subprocess.run(
        [SOFFICE, "-env:UserInstallation=file:///" + TMP_ROOT.replace("\\", "/") + "/lo",
         "--headless", "--convert-to", target, "--outdir", BENCH, src],
        capture_output=True, timeout=180)
    name = os.path.splitext(os.path.basename(src))[0] + "." + target.split(":")[0]
    return os.path.exists(os.path.join(BENCH, name))


def prepare_samples():
    os.makedirs(BENCH, exist_ok=True)
    missing_basic = [n for n in ("sample.txt", "sample.docx") if not os.path.exists(os.path.join(BENCH, n))]
    if missing_basic:
        _make_basic_samples()
    conv = [("sample.docx", "doc:MS Word 97", "sample.doc"),
            ("sample.xlsx", "xls:MS Excel 97", "sample.xls"),
            ("sample.docx", "ppt:MS PowerPoint 97", "sample.ppt"),
            ("sample.docx", "pdf", "sample.pdf")]
    for src, target, dst in conv:
        if not os.path.exists(os.path.join(BENCH, dst)):
            if os.path.exists(os.path.join(BENCH, src)) and os.path.exists(SOFFICE):
                _soffice_convert(os.path.join(BENCH, src), target)
    for name in SAMPLES:
        if not os.path.exists(os.path.join(BENCH, name)):
            SKIP.append(name)


# ---------------- 工具 ----------------

def sha256(b):
    return hashlib.sha256(b).hexdigest()


def pick_version(kind, env_len, cap_codes):
    """按预估码数选最小可用版本（控制回归运行时长；默认口径 v15 优先）。"""
    if it.PROTOCOL_V2:
        # 不可压缩合成帧（x 填充），保证按真实长度估算页数/帧数
        meta = it._env_meta_v2("file", "x", None, env_len)
        probe = it._jz2_frame("env", 0, 1, meta, b"x" * env_len)
    else:
        probe = b"x" * env_len
    for v in (15, 20, 25, 30, 35, 40):
        est = it.estimate_output(kind, "file", "x", probe, v, 1)
        if est.get("pages") is not None and est.get("codes", 0) <= cap_codes:
            return v
    return 40


def roundtrip_static(fmt, name, env_bytes, version):
    """静态二维码全链路还原，返回 (fmt, name, data, ext, 页数)。"""
    out_dir = os.path.join(TMP_ROOT, "static_out")
    os.makedirs(out_dir, exist_ok=True)
    paths, count = it.encode_static_output(fmt, name, env_bytes, version,
                                           out_dir, "t")
    codes = []
    for p in paths:
        with open(p, "rb") as f:
            codes.extend(it.decode_image_file(f.read()))
    env_obj, _info = it._parse_scanned_codes(codes)
    parsed = it.parse_envelope(env_obj)
    return parsed[0], parsed[1], parsed[2], parsed[3], count


def roundtrip_video(fmt, name, env_bytes, version):
    """视频流全链路还原，返回 (fmt, name, data, ext, 帧数)。"""
    out_dir = os.path.join(TMP_ROOT, "video_out")
    os.makedirs(out_dir, exist_ok=True)
    mp4 = os.path.join(out_dir, "t.mp4")
    nframes, _k, _m = it.encode_to_video(env_bytes, version, mp4, frame_repeat=1)
    codes = it.decode_video_frames(mp4)
    raw = it.reassemble_any(codes)
    env_obj = it._envelope_from_raw(raw)
    parsed = it.parse_envelope(env_obj)
    return parsed[0], parsed[1], parsed[2], parsed[3], nframes


def canonical_rows(rows):
    return json.dumps(rows, ensure_ascii=False, sort_keys=True)


def check(label, cond, detail=""):
    RESULTS.append((label, bool(cond), detail))
    print(("PASS " if cond else "FAIL ") + label + (f"  {detail}" if detail and not cond else ""))


# ---------------- 用例 ----------------

def run():
    prepare_samples()
    print(f"样例目录：{BENCH}")
    if SKIP:
        print(f"跳过（无样例）：{', '.join(SKIP)}")

    for name in SAMPLES:
        if name in SKIP:
            continue
        path = os.path.join(BENCH, name)
        with open(path, "rb") as f:
            fb = f.read()
        short = os.path.basename(name)

        # ---- 原件传输（fmt=file）：导出字节必须与源文件一致 ----
        fmt, rname, data, ext = it.extract_doc_raw(name, fb)
        env = it.build_envelope(fmt, rname, data, ext)
        # T02 断言：可压缩源文件的原件信封应启用压缩标记（v1 zip=1 / v2 comp=zlib|xz，
        # T10 起试压取更小者，文本类源文件 xz 通常胜出）
        if env.startswith(it.JZ2_MAGIC):
            fr = it.parse_frame_jz2(env)
            env_obj = it.parse_envelope_v2(env)
            expect_zip = len(zlib.compress(fb, 9)) < len(fb)
            check(f"{short} 原件压缩标记",
                  (fr["comp"] in (it.JZ2_COMP_ZLIB, it.JZ2_COMP_XZ)) == expect_zip,
                  f"comp={fr['comp']} expect={expect_zip}")
            check(f"{short} 原件信封不大于原样字节",
                  len(env) <= len(fb) + it.JZ2_HEADER_LEN + 16)
        else:
            env_obj = json.loads(env.decode("utf-8"))
            z = len(base64.b64encode(zlib.compress(fb, 9)))
            expect_zip = len(base64.b64encode(fb)) > z
            check(f"{short} 原件压缩标记", (env_obj.get("zip") == 1) == expect_zip,
                  f"zip={env_obj.get('zip')} expect={expect_zip}")
            check(f"{short} 原件信封不大于原样 base64",
                  len(env) <= len(json.dumps({"jzt": 1, "fmt": "file", "name": rname,
                                              "data": base64.b64encode(fb).decode()}, ensure_ascii=False)) + 4)

        # 静态（版本按预估页数自适应，上限 120 页控制时长）
        v = pick_version("static", len(env), 120)
        try:
            rfmt, _n, rdata, _e, pages = roundtrip_static(fmt, rname, env, v)
            exported = it.envelope_to_file(rfmt, rname, rdata, _e)
            check(f"{short} 原件·静态还原(v{v},{pages}页)",
                  rfmt == "file" and exported[0] == fb,
                  f"sha {sha256(exported[0])[:12]} vs {sha256(fb)[:12]}")
        except Exception as e:
            check(f"{short} 原件·静态还原", False, f"{type(e).__name__}: {e}")

        # 视频（版本按预估帧数自适应，上限 320 帧控制时长）
        v = pick_version("video", len(env), 320)
        try:
            rfmt, _n, rdata, _e, frames = roundtrip_video(fmt, rname, env, v)
            exported = it.envelope_to_file(rfmt, rname, rdata, _e)
            check(f"{short} 原件·视频还原(v{v},{frames}帧)",
                  rfmt == "file" and exported[0] == fb,
                  f"sha {sha256(exported[0])[:12]} vs {sha256(fb)[:12]}")
        except Exception as e:
            check(f"{short} 原件·视频还原", False, f"{type(e).__name__}: {e}")

        # ---- 精简传输（支持精简的格式）：还原数据结构一致 ----
        try:
            lfmt, lname, ldata, lext = it.extract_doc_lean(name, fb)
        except it.LeanUnsupported:
            continue  # ppt/pdf 仅原件（已覆盖）
        lenv = it.build_envelope(lfmt, lname, ldata, lext)
        v = pick_version("static", len(lenv), 120)
        try:
            rfmt, _n, rdata, _e, pages = roundtrip_static(lfmt, lname, lenv, v)
            if isinstance(ldata, str):
                ok = rfmt == lfmt and rdata == ldata
                detail = "" if ok else "文本不一致"
            else:
                ok = rfmt == lfmt and canonical_rows(rdata) == canonical_rows(ldata)
                detail = "" if ok else "二维数组不一致"
            check(f"{short} 精简·静态还原(v{v},{pages}页)", ok, detail)
        except Exception as e:
            check(f"{short} 精简·静态还原", False, f"{type(e).__name__}: {e}")

    # ---- T02/T04 行为断言 ----
    # fmt=file 不可压缩的小文件不启用压缩标记（信封形状与旧版一致）
    tiny = it.build_envelope("file", "a.bin",
                             base64.b64encode(os.urandom(64)).decode(), None)
    if tiny.startswith(it.JZ2_MAGIC):
        check("fmt=file 不可压缩时保持原形态",
              it.parse_frame_jz2(tiny)["comp"] == it.JZ2_COMP_NONE)
    else:
        check("fmt=file 不可压缩时保持原形态", b'"zip"' not in tiny)
    # 原件路径确实压缩：可压缩源文件的 env 应显著小于 base64 直装
    with open(os.path.join(BENCH, "sample.txt"), "rb") as f:
        tfb = f.read()
    env_t = it.build_envelope("file", "sample.txt",
                              base64.b64encode(tfb).decode("ascii"), None)
    check("fmt=file zlib 试压生效（txt 10×）", len(env_t) < len(tfb) * 0.5,
          f"env={len(env_t)} raw={len(tfb)}")
    # 静态分页不可行 → estimate 给出引导文案（发码前快速失败）
    big = b"A" * (220 * 1024)  # v15 下超 200 页
    est = it.estimate_output("static", "file", "big.bin",
                             base64.b64encode(big), 15, 5)
    check("静态超 200 页前置引导", est.get("pages") is None and "视频" in (est.get("error") or ""))
    # 视频估算口径：秒数与帧数一致
    est = it.estimate_output("video", "file", "x.bin",
                             base64.b64encode(b"B" * 50_000).decode("ascii"), 15, 5)
    check("视频估算帧数×0.333s", est.get("codes", 0) > 0 and
          abs(est["seconds"] - est["codes"] * 5 / 15) < 0.2)

    fails = [r for r in RESULTS if not r[1]]
    print(f"\n==== {len(RESULTS) - len(fails)}/{len(RESULTS)} 通过 ====")
    return 1 if fails else 0


if __name__ == "__main__":
    try:
        code = run()
    finally:
        shutil.rmtree(TMP_ROOT, ignore_errors=True)
    sys.exit(code)
