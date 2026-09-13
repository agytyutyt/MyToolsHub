"""知识库 —— Excel(.xlsx) → 只读 HTML 表格渲染（Python 手绘，样式还原）。

背景
----
xlsx 走 LibreOffice→PDF 预览时按打印分页，宽表/长表被切成多页，不利阅读。
本模块把 xlsx 直接渲染为**连续的单页 HTML 表格**（替代 Excel 类的 PDF 预览，
《插件库优化方案》阶段 3）：一次读全簿，每个 sheet 生成一段
`<table>` HTML 片段（含 <colgroup> 列宽），以 JSON 落盘 `<id>.json`，
前端经 /preview 端点按页签展示。纯只读：不生成任何可编辑控件。

技术选型（参考 GitHub: Apkawa/xlsx2html 的 openpyxl→HTML 思路，按本项目约束重写）
--------
- **零新增 pip 依赖**：只用 openpyxl（旧版格式转换已引入，B-4 优雅降级同口径），
  不引入 xlsx2html 及其传递依赖 babel；数字/日期格式化自实现（覆盖常用格式）。
- 修正 xlsx2html 的几处失真：主题色（theme+tint）它直接丢弃、这里解析
  wb.loaded_theme 还原；列宽换算它用 width/10*96 偏窄，这里用 Excel 标准
  `px ≈ round(width*7)+5`；字号它把 11pt 当 11px，这里按 1pt=4/3px 换算；
  表格用 `table-layout:fixed` 锁定列宽，不被长文本撑变形；
  数字格式用 Excel 的四舍五入（ROUND_HALF_UP，非 Python 默认的银行家舍入）。
- **还原范围**（"尽可能还原原样式"的落地口径）：合并单元格（colspan/rowspan，
  外边框取边缘单元格——即「openpyxl 丢合并边框」的正确读法）、列宽/行高、
  隐藏行列、边框（13 种样式映射）、填充色（solid 取 fgColor；图案填充取底色
  近似）、字体（名称/字号/加粗/斜体/下划线/删除线/颜色，含 theme+tint）、
  对齐（水平/垂直/自动换行/缩进，垂直默认 bottom 对齐 Excel）、
  数字与日期显示格式（千分位/百分比/货币符号/负数括号与 [Red]/中文日期
  'yyyy"年"m"月"d"日"'/时分秒/AM-PM/[h] 经时/科学计数）、网格线开关。
- **不还原**：单元格批注、浮动图片/图表、条件格式、冻结窗格（预览为连续
  长表，无需冻结）、富文本局部样式（openpyxl 默认读为纯文本拼接）。
- **上限熔断**（《插件库优化方案》§3.2 降级链）：行/列/单元格数/单表输出
  字节超限即截断并在表尾提示"下载原件查看全部"；源文件 >15MB 或解析失败
  → 抛 `XlsxRenderError`，记录置 failed，前端回退 SheetJS 简化渲染，
  上传/阅读/下载不受影响（规范 B-4）。
- **安全（SEC-5）**：单元格文本一律转义后输出，样式值仅由本模块拼装
  （颜色经过正则校验），不回传文件内部路径/堆栈；前端注入前仍过
  DOMPurify（纵深防御）。

输出 JSON 结构
--------------
{"version": 1, "generated_at": "...", "truncated": bool,
 "sheets": [{"name": "...", "rows_total": int, "cols_total": int,
             "truncated": bool, "html": "<table ...>...</table>"}]}
"""

import colorsys
import datetime as _dt
import json
import os
import re
import tempfile
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation

try:
    import openpyxl
    from openpyxl.styles.colors import COLOR_INDEX
except Exception:  # pragma: no cover —— 缺依赖时插件其余功能不受影响（B-4）
    openpyxl = None


class XlsxRenderError(Exception):
    """渲染失败，携带**面向用户**的提示文案（不回传堆栈/路径，SEC-5）。"""


# ===================== 渲染上限（防超大表拖垮服务/浏览器） =====================

# 源文件大小上限：openpyxl 常规模式全量载入内存，15MB xlsx 实测可达 GB 级内存，
# 超限直接置 failed 回退前端 SheetJS（其按流解析，浏览器侧更扛得住）。
MAX_SOURCE_BYTES = 15 * 1024 * 1024
MAX_SHEETS = 20          # 最多渲染的 sheet 数
MAX_ROWS = 3000          # 单 sheet 最多行数
MAX_COLS = 120           # 单 sheet 最多列数
MAX_CELLS = 120000       # 单 sheet 最多渲染单元格数
MAX_HTML_BYTES = 4 * 1024 * 1024   # 单 sheet HTML 输出上限（字节）

# Excel 默认列宽（字符数）与像素换算：px ≈ round(width * 7) + 5（Calibri 11 基准，
# 8.43 字符 ≈ 64px，与 Excel 实测一致）
DEFAULT_COL_CHARS = 8.43

_TRUNC_TAIL = "（预览仅显示部分行列，请下载原件查看全部内容）"


# ===================== 颜色 =====================

# 兜底 Office 主题色（theme1.xml 解析失败时使用；与 Excel 2016+ 默认主题一致）
_FALLBACK_SCHEME = {
    "dk1": "000000", "lt1": "FFFFFF", "dk2": "44546A", "lt2": "E7E6E6",
    "accent1": "4472C4", "accent2": "ED7D31", "accent3": "A5A5A5",
    "accent4": "FFC000", "accent5": "5B9BD5", "accent6": "70AD47",
    "hlink": "0563C1", "folHlink": "954F72",
}

_RGB6_RE = re.compile(r"^[0-9A-Fa-f]{6}$")


def _theme_palette(wb):
    """解析主题色表，返回**Excel theme 索引** → "RRGGBB"。

    Excel 的 theme 索引与 clrScheme XML 顺序不同（0~3 交叉）：
    0=lt1（白）1=dk1（黑）2=lt2 3=dk2 4..9=accent1..6 10=hlink 11=folHlink。
    """
    scheme = dict(_FALLBACK_SCHEME)
    try:
        theme = getattr(wb, "loaded_theme", None)
        if isinstance(theme, bytes):
            theme = theme.decode("utf-8", "replace")
        if theme:
            m = re.search(r"<a:clrScheme[ >].*?</a:clrScheme>", theme, re.S)
            if m:
                for name, body in re.findall(
                        r"<a:(dk1|lt1|dk2|lt2|accent[1-6]|hlink|folHlink)>(.*?)</a:\1>",
                        m.group(0), re.S):
                    sm = re.search(r'<a:srgbClr val="([0-9A-Fa-f]{6})"', body)
                    if not sm:
                        sm = re.search(r'<a:sysClr[^>]*lastClr="([0-9A-Fa-f]{6})"', body)
                    if sm:
                        scheme[name] = sm.group(1).upper()
    except Exception:
        pass  # 主题解析失败 → 用兜底色，不阻断渲染
    return [scheme["lt1"], scheme["dk1"], scheme["lt2"], scheme["dk2"],
            scheme["accent1"], scheme["accent2"], scheme["accent3"],
            scheme["accent4"], scheme["accent5"], scheme["accent6"],
            scheme["hlink"], scheme["folHlink"]]


def _tint_rgb(rgb6, tint):
    """按 tint 调整颜色明度（MS 规范的 HLS 明度算法，openpyxl 未实现）。

    tint>0 变亮（向白靠拢），tint<0 变暗（向黑靠拢）。口径与 openpyxl
    官方 issue #1204 给出的实现一致。
    """
    if not tint:
        return rgb6
    try:
        r, g, b = (int(rgb6[i:i + 2], 16) / 255.0 for i in (0, 2, 4))
        h, l, s = colorsys.rgb_to_hls(r, g, b)
        if tint < 0:
            l *= (1.0 + tint)
        else:
            l = l * (1.0 - tint) + tint
        r, g, b = colorsys.hls_to_rgb(h, min(max(l, 0.0), 1.0), s)
        return "".join("%02X" % round(c * 255) for c in (r, g, b))
    except Exception:
        return rgb6


def _color(color, palette):
    """openpyxl Color → "#RRGGBB"；解析不了返回 None（用默认色）。"""
    if color is None:
        return None
    try:
        ctype = color.type
    except Exception:
        return None
    if ctype == "rgb":
        v = color.rgb
        if not isinstance(v, str) or len(v) < 6:
            return None
        v = v[-6:]                      # 取后 6 位（前 2 位是 alpha）
        if not _RGB6_RE.match(v):
            return None
        return "#" + v.upper()
    if ctype == "indexed":
        try:
            v = COLOR_INDEX[color.indexed]
        except Exception:
            return None                 # 64/65 为系统前景/背景色，视作默认
        if not isinstance(v, str) or len(v) < 6:
            return None
        return "#" + v[-6:].upper()
    if ctype == "theme":
        try:
            idx = int(color.theme)
            base = palette[idx] if 0 <= idx < len(palette) else None
        except Exception:
            return None
        if not base:
            return None
        return "#" + _tint_rgb(base, float(color.tint or 0))
    return None                          # auto / 其它 → 默认色


# ===================== 显示值格式化（数字 / 日期时间） =====================

_NUM_COLOR_NAMES = {
    "black": "#000000", "white": "#FFFFFF", "red": "#FF0000",
    "green": "#008000", "blue": "#0000FF", "yellow": "#FFFF00",
    "magenta": "#FF00FF", "cyan": "#00FFFF",
}
# [Red] 等颜色标签 / [>=100] 等条件标签（统一剥掉）
_COLOR_TAG_RE = re.compile(
    r"\[(Red|Green|Blue|White|Magenta|Cyan|Yellow|Black|Color\s*\d+|"
    r">=?\s*[-0-9.]+|<=?\s*[-0-9.]+|=\s*[-0-9.]+)\]", re.I)
_CURRENCY_RE = re.compile(r"\[\$([^\-\]]*)(-[0-9A-Fa-f]+)?\]")
# 数字骨架：至少含一个 #/0（含千分位逗号与小数部分）
_NUMBER_SPAN_RE = re.compile(r"[#0][#0,]*(?:\.[#0]+)?|\.[#0]+")
_QUOTED_RE = re.compile(r'"([^"]*)"')


def _general_number(v):
    """General 格式的数字显示（对齐 Excel：整数不带小数点，浮点去尾零）。"""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, int):
        return str(v) if abs(v) < 1e15 else repr(v)
    f = float(v)
    if f == int(f) and abs(f) < 1e15:
        return str(int(f))
    r = round(f, 10)
    if r == int(r):
        return str(int(r))
    if 0 < abs(r) < 1e-4:
        return ("%.10f" % r).rstrip("0").rstrip(".")
    s = repr(r)
    return s[:-2] if s.endswith(".0") else s


def _fixed(av, dec, frac_optional):
    """定点格式化（Excel 口径 ROUND_HALF_UP，非 Python 默认的银行家舍入）。

    frac_optional：小数位中 `#` 的个数——尾部零要剥掉，但保底 `0` 的位数
    （dec - frac_optional）。返回 (整数部分字符串, 小数部分字符串|None)。
    """
    try:
        q = Decimal(str(av)).quantize(Decimal(1).scaleb(-dec), rounding=ROUND_HALF_UP)
        s = format(q, "f")
    except (InvalidOperation, ValueError):   # pragma: no cover —— 超大数兜底
        s = "%.*f" % (dec, av)
    if "." in s:
        ip, fp = s.split(".", 1)
    else:
        ip, fp = s, ""
    floor_digits = dec - (frac_optional or 0)
    if frac_optional and fp:
        fp = fp.rstrip("0")
        if len(fp) < floor_digits:
            fp = fp.ljust(floor_digits, "0")
        if not fp:
            fp = None
    elif not dec:
        fp = None
    return ip, fp


def _group_digits(digits):
    out = []
    while len(digits) > 3:
        out.insert(0, digits[-3:])
        digits = digits[:-3]
    out.insert(0, digits)
    return ",".join(out)


def _fmt_number(value, fmt):
    """数字 → (显示文本, 颜色覆盖|None)。

    覆盖常用格式：General / 0、0.00 / #,##0(.00) / 百分比 / 货币前缀（¥ $ € £ 等
    字面量与 [$$-409] 标签）/ 负数独立 section（括号、-）与 [Red] / 科学计数 E+00 /
    尾逗号千分缩放。`?` 补位、颜色编号等按近似处理（去占位符，不改变数值主体）。
    """
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), None
    fmt = (fmt or "General").strip() or "General"
    if fmt.lower() == "general":
        return _general_number(value), None
    if "/" in fmt.replace("\\", ""):
        # 分数格式（# ?/? 等）不支持占位还原 → 按 General 数值显示，避免乱码
        return _general_number(value), None

    sections = fmt.split(";")
    v = value if isinstance(value, int) else float(value)
    neg_sep = len(sections) > 1 and _NUMBER_SPAN_RE.search(sections[1])
    sec_idx = 0
    if v < 0:
        if neg_sep:
            sec_idx = 1
        v = -v
    elif v == 0 and len(sections) > 2 and _NUMBER_SPAN_RE.search(sections[2]):
        sec_idx = 2
    sec = sections[min(sec_idx, len(sections) - 1)]

    color_override = None
    m = _COLOR_TAG_RE.search(sec)
    if m and m.group(1).lower() in _NUM_COLOR_NAMES:
        color_override = _NUM_COLOR_NAMES[m.group(1).lower()]
    sec = _CURRENCY_RE.sub(lambda mm: mm.group(1) or "", sec)   # [$$-409] → $
    sec = _COLOR_TAG_RE.sub("", sec)                            # [Red] / 条件 → 去掉
    sec = sec.replace("\\", "")
    sec = re.sub(r"_.", " ", sec)       # _x → 一个空格宽度
    sec = re.sub(r"\*.", "", sec)       # *x → 重复填充符，去掉
    sec = sec.replace("?", "#")         # ? 补位占位 → 按 # 近似

    literals = []

    def _stash(mm):
        literals.append(mm.group(1))
        return "\x00"

    sec = _QUOTED_RE.sub(_stash, sec)

    pct = sec.count("%")
    if pct:
        v = v * (100 ** pct)

    em = re.search(r"[Ee][+-]0+", sec)
    if em:                               # 科学计数：E 后 0 的个数为小数位
        digits = len(re.search(r"0+", em.group(0)).group(0))
        mant = re.sub(_NUMBER_SPAN_RE, "", sec.replace(em.group(0), ""))
        return (mant + "%.*E" % (digits, v)).replace("\x00", ""), color_override

    nm = _NUMBER_SPAN_RE.search(sec)
    if not nm:                           # 无数字骨架（纯字面量格式）：显示字面量
        out = sec
        for lit in literals:
            out = out.replace("\x00", lit, 1)
        return out.replace("\x00", ""), color_override

    prefix = sec[:nm.start()]
    suffix = sec[nm.end():]
    for lit in literals:                 # 按原顺序还原到前/后缀里
        if "\x00" in prefix:
            prefix = prefix.replace("\x00", lit, 1)
        else:
            suffix = suffix.replace("\x00", lit, 1)
    prefix, suffix = prefix.replace("\x00", ""), suffix.replace("\x00", "")

    int_part, _, frac_part = nm.group(0).partition(".")
    # 尾逗号 = 千分缩放（#,##0, → 按千显示）
    mm2 = re.search(r",+$", int_part)
    if mm2:
        v = v / (1000 ** len(mm2.group(0)))
        int_part = int_part[:mm2.start()]
    grouping = "," in int_part
    min_int = int_part.count("0")
    frac_dec = len(frac_part)
    frac_opt = frac_part.count("#")

    ip, fp = _fixed(abs(v), frac_dec, frac_opt)
    if grouping:
        ip = _group_digits(ip)
    if min_int and len(ip) < min_int:
        ip = ip.zfill(min_int)
    txt = prefix + ip + (("." + fp) if fp else "") + suffix
    if value < 0 and not neg_sep:
        txt = "-" + txt
    return txt, color_override


_MONTHS_EN = ["January", "February", "March", "April", "May", "June", "July",
              "August", "September", "October", "November", "December"]
_WEEKDAYS_EN = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
                "Saturday", "Sunday"]

# 日期时间格式 token（长 token 优先；引号/方括号整体吞掉）
_DT_TOKEN_RE = re.compile(
    r'"[^"]*"|\[[^\]]*\]|AM/PM|A/P|mmmmm|mmmm|mmm|dddd|ddd|dd|yyyy|yyy|yy|'
    r"hh|mm|ss|\.0+|y|m|d|h|s|e|.", re.I)


def _fmt_datetime(v, fmt):
    """日期/时间/经时 → 显示文本（支持中文格式 'yyyy"年"m"月"d"日"' 等）。

    m/mm 的月份 vs 分钟消歧按 Excel 规则：紧随 h/hh/[h..] 之后、或其后紧跟
    s/ss/.0 的 m 视为分钟，其余视为月份。
    """
    if isinstance(v, _dt.timedelta):
        total = int(v.total_seconds())
        h, rem = divmod(abs(total), 3600)
        mi, se = divmod(rem, 60)
        fields = {"yyyy": None, "yy": None, "e": None, "y": None,
                  "m": None, "mm": None, "d": None, "dd": None, "dw": None,
                  "h": h, "hh": h, "mi": mi, "s": se, "ss": se}
        elapsed = True
        micros = 0
        hour24 = h
    elif isinstance(v, _dt.datetime):
        fields = {"yyyy": v.year, "yy": v.year % 100, "e": v.year, "y": v.year % 100,
                  "m": v.month, "mm": v.month, "d": v.day, "dd": v.day,
                  "dw": v.weekday(),
                  "h": v.hour % 12 or 12, "hh": v.hour, "mi": v.minute,
                  "s": v.second, "ss": v.second}
        elapsed = False
        micros = v.microsecond
        hour24 = v.hour
    elif isinstance(v, _dt.date):
        fields = {"yyyy": v.year, "yy": v.year % 100, "e": v.year, "y": v.year % 100,
                  "m": v.month, "mm": v.month, "d": v.day, "dd": v.day,
                  "dw": v.weekday(),
                  "h": None, "hh": None, "mi": None, "s": None, "ss": None}
        elapsed = False
        micros = 0
        hour24 = None
    else:  # datetime.time
        fields = {"yyyy": None, "yy": None, "e": None, "y": None,
                  "m": None, "mm": None, "d": None, "dd": None, "dw": None,
                  "h": v.hour % 12 or 12, "hh": v.hour, "mi": v.minute,
                  "s": v.second, "ss": v.second}
        elapsed = False
        micros = v.microsecond
        hour24 = v.hour

    fmt = (fmt or "").split(";")[0] or "General"
    if fmt.lower() == "general":
        if isinstance(v, _dt.datetime):
            return v.strftime("%Y-%m-%d %H:%M:%S")
        if isinstance(v, _dt.date):
            return v.strftime("%Y-%m-%d")
        if isinstance(v, _dt.time):
            return v.strftime("%H:%M:%S")
        return str(v)

    sec = fmt.replace("\\", "")
    sec = re.sub(r"_.", " ", sec)
    sec = re.sub(r"\*.", "", sec)
    sec = _COLOR_TAG_RE.sub("", sec)

    tokens = _DT_TOKEN_RE.findall(sec)
    out = []
    prev_h = False

    def _next_meaningful(idx):
        """m/mm 之后第一个**日期时间相关** token（跳过引号字面量与普通字面量），
        用于分钟/月份消歧（Excel 规则：m 紧随 h 或后跟 s 即为分钟）。"""
        for t in tokens[idx + 1:]:
            if t.startswith('"') or t.startswith("["):
                continue
            lt = t.lower()
            if re.match(r"^[ymdhs.0]+$", lt):
                return lt
        return ""

    for i, tok in enumerate(tokens):
        low = tok.lower()
        if tok.startswith('"'):
            out.append(tok[1:-1])
        elif tok.startswith("["):
            inner = tok[1:-1].lower()
            if elapsed and re.match(r"^h+$", inner):
                out.append(str(fields["h"]))       # [h] 经时：累计小时
                prev_h = True
            # 其余 [..]（语言标签等）丢弃
        elif low == "am/pm":
            out.append(("AM" if hour24 is not None and hour24 < 12 else "PM")
                       if hour24 is not None else "")
            prev_h = False
        elif low == "a/p":
            out.append(("A" if hour24 is not None and hour24 < 12 else "P")
                       if hour24 is not None else "")
            prev_h = False
        elif low in ("yyyy", "yyy", "yy", "y", "e"):
            y = fields.get("yyyy" if low in ("yyyy", "yyy", "e") else "yy")
            out.append("" if y is None else
                       (str(y) if low in ("yyyy", "e") else "%02d" % (y % 100)))
            prev_h = False
        elif low in ("m", "mm"):
            if elapsed:
                val = fields["mi"]      # 经时格式无月份概念，m 即分钟
                out.append("" if val is None else
                           ("%02d" % val if low == "mm" else str(val)))
            elif fields["m"] is None and fields["mi"] is None:
                out.append("")
            else:
                minute_ctx = prev_h or _next_meaningful(i) in ("s", "ss") \
                    or _next_meaningful(i).startswith(".0")
                if minute_ctx:
                    val = fields["mi"]
                    out.append("" if val is None else
                               ("%02d" % val if low == "mm" else str(val)))
                else:
                    mo = fields["m"]
                    if mo is None:
                        out.append("")
                    elif low == "mm":
                        out.append("%02d" % mo)
                    else:
                        out.append(str(mo))
            prev_h = False
        elif low in ("mmmmm", "mmmm", "mmm"):
            mo = fields["m"]
            if elapsed or mo is None:
                out.append("")
            elif low == "mmmmm":
                out.append(_MONTHS_EN[mo - 1][0])
            elif low == "mmmm":
                out.append(_MONTHS_EN[mo - 1])
            else:
                out.append(_MONTHS_EN[mo - 1][:3])
            prev_h = False
        elif low in ("d", "dd", "ddd", "dddd"):
            if elapsed or (low in ("ddd", "dddd") and fields["dw"] is None)                     or (low in ("d", "dd") and fields["d"] is None):
                out.append("")
            elif low == "dddd":
                out.append(_WEEKDAYS_EN[fields["dw"]])
            elif low == "ddd":
                out.append(_WEEKDAYS_EN[fields["dw"]][:3])
            else:
                out.append("%02d" % fields["d"] if low == "dd" else str(fields["d"]))
            prev_h = False
        elif low in ("h", "hh"):
            hv = fields.get(low)
            if hv is None:
                out.append("")
            elif elapsed:
                out.append(str(hv))
            elif low == "hh":
                out.append("%02d" % (hv % 24))
            else:
                out.append(str(hv % 24))
            prev_h = True
        elif low in ("s", "ss"):
            sv = fields.get(low)
            out.append("" if sv is None else
                       ("%02d" % sv if low == "ss" else str(sv)))
            prev_h = False
        elif re.match(r"^\.0+$", tok):           # 秒的小数（ss.0 / ss.00）
            digits = len(tok) - 1
            out.append("." + (("%06d" % micros) if micros else "0" * 6)[:digits])
            prev_h = False
        else:
            out.append(tok)
            # 普通字面量（":"、"年" 等）不打断 h→mm 的分钟上下文（对齐 Excel）
    return "".join(out)


def _display(cell):
    """单元格 → (显示文本, 颜色覆盖|None)。文本未转义（由调用方统一转义）。"""
    v = cell.value
    if v is None:
        return "", None
    fmt = getattr(cell, "number_format", None) or "General"
    if isinstance(v, bool):
        return ("TRUE" if v else "FALSE"), None
    if isinstance(v, str):
        return v, None
    if isinstance(v, ( _dt.datetime, _dt.date, _dt.time, _dt.timedelta)):
        return _fmt_datetime(v, fmt), None
    if isinstance(v, (int, float)):
        return _fmt_number(v, fmt)
    return str(v), None


# ===================== 样式 → CSS =====================

# openpyxl 边框样式 → CSS (宽度, 线型)
_BORDER_MAP = {
    "thin": ("1px", "solid"), "medium": ("2px", "solid"), "thick": ("3px", "solid"),
    "double": ("3px", "double"), "hair": ("1px", "solid"), "dotted": ("1px", "dotted"),
    "dashed": ("1px", "dashed"), "dashDot": ("1px", "dashed"),
    "dashDotDot": ("1px", "dashed"), "mediumDashed": ("2px", "dashed"),
    "mediumDashDot": ("2px", "dashed"), "mediumDashDotDot": ("2px", "dashed"),
    "slantDashDot": ("2px", "dashed"),
}

_FONT_FALLBACK = "Calibri, 'Microsoft YaHei', sans-serif"


def _border_side(border_obj, side, palette):
    """取某一边框 → "2px solid #FF0000" | None。"""
    b = getattr(border_obj, side, None)
    if b is None or not b.style:
        return None
    mapped = _BORDER_MAP.get(b.style) or ("1px", "solid")   # 未知样式按细实线近似
    color = _color(b.color, palette) or "#000000"
    return "%s %s %s" % (mapped[0], mapped[1], color)


class _Span(object):
    """轻量合并区描述（只为边框查询，避免构造 openpyxl 对象）。"""

    def __init__(self, min_row, min_col, rowspan, colspan):
        self.min_row = min_row
        self.min_col = min_col
        self.max_row = min_row + rowspan - 1
        self.max_col = min_col + colspan - 1


def _merged_edge_borders(ws, span, palette):
    """合并区外边框：每侧取该侧边缘单元格的第一条有效边框。

    openpyxl 合并后只有锚点保留值，边框记录在各边缘单元格上——这正是
    「openpyxl 丢合并边框」问题的根源，这里按边缘取值还原。
    """
    res = {}
    for side, cells in (
            ("top", [(span.min_row, c) for c in range(span.min_col, span.max_col + 1)]),
            ("bottom", [(span.max_row, c) for c in range(span.min_col, span.max_col + 1)]),
            ("left", [(r, span.min_col) for r in range(span.min_row, span.max_row + 1)]),
            ("right", [(r, span.max_col) for r in range(span.min_row, span.max_row + 1)])):
        for r, c in cells:
            css = _border_side(ws.cell(row=r, column=c).border, side, palette)
            if css:
                res[side] = css
                break
    return res


def _cell_css(cell, palette, extra_borders=None):
    """单元格样式 → inline CSS（只输出与默认不同的项，控制体积）。"""
    css = []
    f = cell.font
    if f is not None:
        if f.name and str(f.name) != "Calibri":
            css.append("font-family:'%s',%s" % (str(f.name).replace("'", ""), _FONT_FALLBACK))
        try:
            sz = float(f.sz or 11)
        except (TypeError, ValueError):
            sz = 11.0
        if abs(sz - 11.0) > 1e-9:
            css.append("font-size:%spx" % round(sz * 4 / 3, 1))
        if f.b:
            css.append("font-weight:bold")
        if f.i:
            css.append("font-style:italic")
        deco = []
        if f.u:
            deco.append("underline double" if str(f.u) == "double" else "underline")
        if f.strike:
            deco.append("line-through")
        if deco:
            css.append("text-decoration:" + " ".join(deco))
        col = _color(f.color, palette)
        if col and col.upper() != "#000000":    # 黑色为默认，省略
            css.append("color:" + col)

    a = cell.alignment
    if a is not None:
        ha = (a.horizontal or "").lower()
        if ha in ("left", "right", "center"):
            css.append("text-align:" + ha)
        va = (a.vertical or "").lower()
        css.append("vertical-align:" +
                   ("middle" if va == "center" else va if va in ("top", "middle", "bottom")
                    else "bottom"))          # Excel 默认垂直靠下
        if a.wrap_text:
            css.append("white-space:pre-wrap")
        else:
            css.append("white-space:nowrap;overflow:hidden")
        if a.indent:
            try:
                css.append("padding-left:%dpx" % (int(a.indent) * 9))
            except (TypeError, ValueError):
                pass
    else:
        css.append("vertical-align:bottom;white-space:nowrap;overflow:hidden")

    fill = cell.fill
    if fill is not None and fill.patternType:
        if fill.patternType == "solid":
            bg = _color(fill.fgColor, palette)
        else:  # 图案填充取底色近似（点阵无法还原，取占主导的背景色）
            bg = _color(fill.bgColor, palette) or _color(fill.fgColor, palette)
        if bg and bg.upper() != "#FFFFFF":  # 白底为默认，省略
            css.append("background-color:" + bg)

    borders = extra_borders or {}
    border = cell.border
    for side in ("top", "right", "bottom", "left"):
        css_side = borders.get(side) or _border_side(border, side, palette)
        if css_side:
            css.append("border-%s:%s" % (side, css_side))
    return ";".join(css)


# ===================== 工作表 → HTML =====================

def _col_widths(ws, n_cols):
    """列宽数组与隐藏标记（下标 0 ↔ 第 1 列）。"""
    dims = []
    for _letter, dim in list(ws.column_dimensions.items()):
        try:
            if dim.min and dim.max:
                dims.append((int(dim.min), int(dim.max), dim))
        except Exception:
            continue
    default_chars = getattr(ws.sheet_format, "defaultColWidth", None)
    try:
        default_chars = float(default_chars) if default_chars else DEFAULT_COL_CHARS
    except (TypeError, ValueError):
        default_chars = DEFAULT_COL_CHARS
    default_px = round(default_chars * 7) + 5
    widths, hidden = [], []
    for ci in range(1, n_cols + 1):
        dim = next((d for lo, hi, d in dims if lo <= ci <= hi), None)
        w = None
        if dim is not None:
            try:
                if dim.customWidth and dim.width:
                    w = round(float(dim.width) * 7) + 5
            except (TypeError, ValueError):
                w = None
        widths.append(w or default_px)
        hidden.append(bool(dim is not None and dim.hidden))
    return widths, hidden


def _sheet_html(ws, palette):
    """单个 sheet → (html, rows_total, cols_total, truncated)。"""
    max_row = int(ws.max_row or 0)
    max_col = int(ws.max_column or 0)
    n_rows = min(max_row, MAX_ROWS)
    n_cols = min(max_col, MAX_COLS)
    truncated = max_row > MAX_ROWS or max_col > MAX_COLS

    if n_rows <= 0 or n_cols <= 0:      # 空表
        return ('<table class="kb-xlsx-table" '
                'style="table-layout:fixed;border-collapse:collapse"></table>',
                max_row, max_col, False)

    anchors = {}     # (r, c) 锚点 → (rowspan, colspan)
    covered = set()  # 被合并覆盖（非锚点）的格子
    try:
        ranges = list(ws.merged_cells.ranges)
    except Exception:
        ranges = []
    for rng in ranges:
        anchor = (rng.min_row, rng.min_col)
        anchors[anchor] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != anchor:
                    covered.add((r, c))

    widths, col_hidden = _col_widths(ws, n_cols)

    grid = ' class="kb-xlsx-table%s"' % (
        " kb-nogrid" if not getattr(ws.sheet_view, "showGridLines", True) else "")
    parts = ['<table%s style="table-layout:fixed;border-collapse:collapse">'
             "<colgroup>" % grid]
    for ci in range(n_cols):
        if col_hidden[ci]:
            parts.append('<col style="display:none;width:0"/>')
        else:
            parts.append('<col style="width:%dpx"/>' % widths[ci])
    parts.append("</colgroup>")

    cells_done = 0
    bytes_done = 0
    stop = False
    for row in ws.iter_rows(min_row=1, max_row=n_rows, max_col=n_cols):
        r = row[0].row
        dim = ws.row_dimensions.get(r)
        if dim is not None and dim.hidden:
            continue
        tr_attrs = ""
        if dim is not None and dim.customHeight and dim.height:
            try:
                tr_attrs = ' style="height:%dpx"' % round(float(dim.height) * 4 / 3)
            except (TypeError, ValueError):
                pass
        parts.append("<tr%s>" % tr_attrs)
        for cell in row:
            if col_hidden[cell.column - 1]:
                continue
            pos = (r, cell.column)
            if pos in covered:
                continue
            span = anchors.get(pos)
            attrs, extra = "", None
            if span:
                extra = _merged_edge_borders(
                    ws, _Span(r, cell.column, span[0], span[1]), palette)
                if span[0] > 1:
                    attrs += ' rowspan="%d"' % span[0]
                if span[1] > 1:
                    attrs += ' colspan="%d"' % span[1]
            text, color_ovr = _display(cell)
            css = _cell_css(cell, palette, extra_borders=extra)
            if color_ovr:                       # [Red] 等格式色覆盖字体色
                css = (css + ";color:" + color_ovr) if css else ("color:" + color_ovr)
            esc = (text.replace("&", "&amp;").replace("<", "&lt;")
                   .replace(">", "&gt;").replace('"', "&quot;"))
            parts.append('<td%s style="%s">%s</td>' % (attrs, css, esc))
            cells_done += 1
            bytes_done += len(parts[-1])
            if cells_done >= MAX_CELLS or bytes_done >= MAX_HTML_BYTES:
                truncated = True
                stop = True
                break
        parts.append("</tr>")
        if stop:
            break

    if truncated:
        parts.append('<tr class="kb-xlsx-more"><td colspan="%d" '
                     'style="text-align:center;background:#fffbe6;color:#8a6d3b;'
                     'white-space:normal">%s</td></tr>' % (max(n_cols, 1), _TRUNC_TAIL))
    parts.append("</table>")
    return "".join(parts), max_row, max_col, truncated


# ===================== 入口 =====================

def availability():
    """渲染依赖可用性（B-4：缺 openpyxl 时优雅降级）。"""
    return {"available": openpyxl is not None,
            "error": None if openpyxl is not None else "服务器缺少 openpyxl 依赖"}


def render_xlsx(src_path, dst_json):
    """把 xlsx 渲染为 JSON（各 sheet 的 HTML 片段），**原子落盘**到 dst_json。

    成功返回摘要 dict（sheet 数 / 是否截断）；失败抛 `XlsxRenderError`
    （文案可直接展示给用户，SEC-5）。
    """
    if openpyxl is None:
        raise XlsxRenderError("服务器缺少 openpyxl 依赖，无法生成表格样式预览")
    if not src_path or not os.path.isfile(src_path):
        raise XlsxRenderError("源文件不存在，无法生成表格样式预览")
    try:
        if os.path.getsize(src_path) > MAX_SOURCE_BYTES:
            raise XlsxRenderError("文件过大，未能生成表格样式预览，已回退简化渲染")
    except OSError:
        raise XlsxRenderError("源文件不可读，无法生成表格样式预览")

    try:
        wb = openpyxl.load_workbook(src_path, data_only=True)
    except Exception:
        raise XlsxRenderError("表格解析失败，已回退简化渲染（可下载原件查看）")

    try:
        palette = _theme_palette(wb)
        sheets = []
        for name in wb.sheetnames[:MAX_SHEETS]:
            try:
                ws = wb[name]
            except Exception:
                continue
            html, rows_total, cols_total, truncated = _sheet_html(ws, palette)
            sheets.append({"name": name, "rows_total": rows_total,
                           "cols_total": cols_total, "truncated": truncated,
                           "html": html})
        if not sheets:
            raise XlsxRenderError("表格无可用工作表，已回退简化渲染")
        payload = {
            "version": 1,
            "generated_at": _dt.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "truncated": any(s["truncated"] for s in sheets),
            "sheets": sheets,
        }
        try:
            dst_dir = os.path.dirname(os.path.abspath(dst_json))
            os.makedirs(dst_dir, exist_ok=True)
            fd, tmp = tempfile.mkstemp(prefix=".xlsxhtml-", dir=dst_dir)
            with os.fdopen(fd, "w", encoding="utf-8") as f:
                json.dump(payload, f, ensure_ascii=False)
            os.replace(tmp, dst_json)   # 同分区原子替换
        except OSError:
            raise XlsxRenderError("预览数据保存失败，请稍后重试")
        return {"sheets": len(sheets), "truncated": payload["truncated"]}
    finally:
        try:
            wb.close()
        except Exception:
            pass
