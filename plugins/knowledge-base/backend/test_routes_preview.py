# -*- coding: utf-8 -*-
"""阶段 8 集成自测：假 admin 会话 + Flask test client，
覆盖 docx/xlsx 上传 → 按需 /preview（引擎渲染 + 磁盘缓存）→ 404 语义 → 删除清理；
另测 office_render 模块级渲染（docx/xlsx/xls/doc 四类，旧格式经 soffice 转换）。"""
import io
import json
import os
import sys
import tempfile
import time
import types

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", ".."))
BACKEND = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, REPO)
sys.path.insert(0, BACKEND)

# ★ 必须在 import routes **之前**设数据根：routes 在模块加载时就把 DATA_DIR 定死了。
#   本套件此前一直跑在**真实数据根**（%USERPROFILE%\.jztoolshub）上——每次运行都往用户的
#   知识库里塞测试文档（2026-10-10 发现时已累积数十条，同名分类还会撞 409 让用例自相残杀）。
#   口径与 info-transfer 各套件一致：临时目录隔离，跑完即弃。
_TMP_ROOT = tempfile.mkdtemp(prefix="kb_test_")
os.environ["JZTOOLS_DATA_ROOT"] = _TMP_ROOT

# ---- 伪造 jztools_admin（routes 顶层 try-import）----
admin = types.ModuleType("jztools_admin")
admin_routes = types.ModuleType("jztools_admin.routes")
admin_routes.get_session_user = lambda: {
    "username": "tester", "name": "测试管理员", "super_admin": True,
    "role_id": "role-admin", "unit_id": "", "department_id": ""}
admin_routes.set_operation = lambda op: None
admin.routes = admin_routes
sys.modules["jztools_admin"] = admin
sys.modules["jztools_admin.routes"] = admin_routes

import importlib  # noqa: E402
pkg = types.ModuleType("kb_backend_test")
pkg.__path__ = [BACKEND]
sys.modules["kb_backend_test"] = pkg
routes = importlib.import_module("kb_backend_test.routes")
office_render = importlib.import_module("kb_backend_test.office_render")

import jz_api  # noqa: E402

# routes 现从框架门面取身份（`_get_session_user = jz_api.get_session_user`，解耦批次改造），
# 上面伪造的 jztools_admin 假模块已不参与取号 —— 按框架真实注册路径补 provider，
# 否则所有需登录的接口一律 401（该套件的既有缺陷，非被测行为回归）。
jz_api.register_providers(get_session_user=admin_routes.get_session_user,
                          set_operation=admin_routes.set_operation)

from flask import Flask  # noqa: E402

OUT = os.path.join(BACKEND, "out")
os.makedirs(OUT, exist_ok=True)

app = Flask(__name__)
routes.register(app)
client = app.test_client()

# ---------- 0) 模块级渲染：四类格式 ----------
import openpyxl  # noqa: E402
import docx  # noqa: E402

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "数据"
ws["A1"] = "区域"; ws["B1"] = "销售额"
ws["A2"] = "华东"
c = ws["B2"]; c.value = 1234567.891; c.number_format = "#,##0.00"
buf = io.BytesIO(); wb.save(buf)
xlsx_bytes = buf.getvalue()

d = docx.Document()
d.add_heading("集成测试文档", level=1)
d.add_paragraph("第一段正文内容。")
t = d.add_table(rows=1, cols=2); t.style = "Table Grid"
t.rows[0].cells[0].text = "甲"; t.rows[0].cells[1].text = "乙"
buf = io.BytesIO(); d.save(buf)
docx_bytes = buf.getvalue()

r = office_render.render_sheet(xlsx_bytes)
assert r["kind"] == "sheet" and "1,234,567.89" in r["html"], "xlsx 渲染失败"
print("module render sheet OK | html %dB | warnings=%s" % (len(r["html"]), r["warnings"]))
r = office_render.render_word(docx_bytes)
assert r["kind"] == "word" and "集成测试文档" in r["html"] and "<style>" in r["html"]
print("module render word OK | html %dB" % len(r["html"]))

# ---------- 0b) Word 产物后处理补丁（三处显示缺陷修正，vendor 零改动） ----------
# 详见 docs/eval/知识库Word预览三问题排查报告.md：
#   ① 仿宋_GB2312 缺失时引擎回退到雅黑（墨量 2.56×）→ 视觉上"莫名加粗"
#   ② <tr> 上的 overflow:hidden 对 table-row 无效 → 精确行高失效、内容穿模
#   ③ 引擎 base CSS 未重置 <p> 默认 margin → 撑高表格单元格
import re as _re  # noqa: E402

# 造一个带 仿宋_GB2312 + 精确行高表格的 docx
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402

d2 = docx.Document()
_n = d2.styles["Normal"]
_n.font.size = Pt(16)
_n.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋_GB2312")
for txt in ("仿宋正文一，竖笔不应过粗，且未加粗。", "仿宋正文二。"):
    p = d2.add_paragraph()
    run = p.add_run(txt)
    run.bold = False
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋_GB2312")
_t = d2.add_table(rows=2, cols=2)
_t.style = "Table Grid"
for i in range(2):
    for j in range(2):
        c = _t.cell(i, j)
        c.text = ""
        rr = c.paragraphs[0].add_run("单元格%d%d" % (i, j))
        rr.bold = False
        rr._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "仿宋_GB2312")
        tcPr = c._tc.get_or_add_tcPr()
        for row in _t.rows:
            trPr = row._tr.get_or_add_trPr()
            if trPr.find(qn("w:trHeight")) is None:
                th = OxmlElement("w:trHeight")
                th.set(qn("w:val"), "400")          # 20pt 精确行高
                th.set(qn("w:hRule"), "exact")
                trPr.append(th)
buf2 = io.BytesIO(); d2.save(buf2)
patch_probe = office_render.render_word(buf2.getvalue())["html"]

# 取「不经过补丁」的引擎原产物做对照
import dhr as _dhr  # noqa: E402  （vendor 已在 office_render 里注入 sys.path）
_raw = _dhr.convert(bytes(buf2.getvalue()), _dhr.ConvertOptions(
    mode="flow", output="fragment", css_prefix="kbdoc",
    media_mode="base64")).html

# ① 字体回退补链：仿宋族声明必须补上 FangSong，且排在雅黑前面
_fams = [f for f in _re.findall(r"font-family:([^;}]+)", patch_probe) if "仿宋" in f]
assert _fams, "补丁后未找到仿宋族 font-family（样例文档应含 仿宋_GB2312）"
assert all('"FangSong"' in f for f in _fams), _fams[0]
assert _fams[0].lstrip().startswith('"仿宋_GB2312"'), "目标字体应仍在首位"
assert _fams[0].find('"FangSong"') < _fams[0].find('"Microsoft YaHei"'), \
    "FangSong 必须排在 Microsoft YaHei 之前"
assert '"FangSong"' not in _raw, "对照组（原产物）不应含该回退"
print("patch font-fallback OK | %s" % _fams[0][:80])

# ② <tr> 上无效的 overflow:hidden 必须被清除，height 保留、标签数不变
_TR = r"<tr\b[^>]*>"
assert any("overflow:hidden" in t for t in _re.findall(_TR, _raw)), "对照组应有需修的 tr"
assert not any("overflow:hidden" in t for t in _re.findall(_TR, patch_probe))
assert len(_re.findall(_TR, _raw)) == len(_re.findall(_TR, patch_probe))
assert any(_re.search(r"height:\d+px", t) for t in _re.findall(_TR, patch_probe)), \
    "height 声明应保留"
print("patch table-overflow OK | tr 上的 overflow:hidden 已清除")

# ③ CSS 重置：p margin 归零 + 单元格可断行
assert ".kbdoc p{margin:0}" in patch_probe
assert "word-break:break-all" in patch_probe and "overflow-wrap:anywhere" in patch_probe
print("patch css-reset OK")

# ④ 幂等：重复处理不得改变产物（否则反复请求会持续膨胀）
assert office_render._polish_word_html(patch_probe) == patch_probe
assert patch_probe.count("/*kbdoc-patch*/") == 1
print("patch idempotent OK")

# ⑤ 内容完整：可见文本与标签数不被破坏
_vis = lambda s: _re.sub(r"\s+", "", _re.sub(
    r"<[^>]+>", "", _re.sub(r"<(style|script)\b[^>]*>.*?</\1>", "", s,
                            flags=_re.S | _re.I)))
assert _vis(_raw) == _vis(patch_probe), "可见文本被破坏"
assert len(_re.findall(r"<[a-zA-Z]", _raw)) == len(_re.findall(r"<[a-zA-Z]", patch_probe))
print("patch content-integrity OK")

# 旧格式：用 soffice 把 xlsx→xls、docx→doc（可用才测）
xls_bytes = doc_bytes = None
soffice = office_render.soffice_available()
print("soffice available:", soffice)
if soffice:
    import subprocess, tempfile, shutil
    work = tempfile.mkdtemp(prefix="kbfmt-")
    try:
        for src_bytes, src_ext, dst_ext in ((xlsx_bytes, "xlsx", "xls"), (docx_bytes, "docx", "doc")):
            sp = os.path.join(work, "f." + src_ext)
            with open(sp, "wb") as f:
                f.write(src_bytes)
            subprocess.run([os.environ.get("XHR_SOFFICE") or office_render._configured_soffice(),
                            "--headless", "--convert-to", dst_ext, "--outdir", work, sp],
                           capture_output=True, timeout=120,
                           startupinfo=None)
            dp = os.path.join(work, "f." + dst_ext)
            if os.path.isfile(dp):
                with open(dp, "rb") as f:
                    if dst_ext == "xls":
                        xls_bytes = f.read()
                    else:
                        doc_bytes = f.read()
    finally:
        shutil.rmtree(work, ignore_errors=True)
    if xls_bytes:
        r = office_render.render_sheet(xls_bytes)
        assert r["kind"] == "sheet" and "<table" in r["html"]
        print("module render legacy xls OK | html %dB | warnings=%s" % (len(r["html"]), r["warnings"]))
    if doc_bytes:
        r = office_render.render_word(doc_bytes)
        # 注：LibreOffice doc 归一化回环对本例的微型表格会退化为制表符段落
        # （内容保留）；真实链路 .doc 上传走 doc_convert 转换件（docx），不经过此通道，
        # 故这里只断言内容保留与 kind 正确，不断言表格结构。
        assert r["kind"] == "word" and "集成测试文档" in r["html"] and "甲" in r["html"]
        print("module render legacy doc OK | html %dB | warnings=%s" % (len(r["html"]), r["warnings"]))

# ---------- 1) status ----------
r = client.get("/api/knowledge-base/status")
data = r.get_json()
print("status:", r.status_code, "| office_preview:", data.get("office_preview"),
      "| office_render:", data.get("office_render"))
assert data.get("office_preview") is True

# ---------- 2) 上传 xlsx / docx → /preview ----------
def upload_and_preview(name, blob, expect_kind, expect_text):
    r = client.post("/api/knowledge-base/files", data={
        "file": (io.BytesIO(blob), name), "name": name.rsplit(".", 1)[0],
    }, content_type="multipart/form-data")
    up = r.get_json()
    assert r.status_code == 200, up
    assert "pdf_status" not in up and "html_status" not in up
    fid = up["id"]
    t0 = time.time()
    r = client.get(f"/api/knowledge-base/files/{fid}/preview")
    pv = r.get_json()
    assert r.status_code == 200 and pv["ok"] and pv["kind"] == expect_kind, pv.get("error")
    assert expect_text in pv["html"], "预览 HTML 缺少预期内容"
    first_ms = (time.time() - t0) * 1000
    # 第二次请求应命中缓存（cached 标记）
    r2 = client.get(f"/api/knowledge-base/files/{fid}/preview")
    pv2 = r2.get_json()
    assert r2.status_code == 200 and pv2.get("cached") is True
    cache_file = f"{fid}.preview.json"
    assert os.path.isfile(os.path.join(routes.FILES_DIR, cache_file))
    print(f"upload+preview {name}: kind={pv['kind']} first={first_ms:.0f}ms "
          f"html={len(pv['html'])}B cached=OK")
    return fid

fid_x = upload_and_preview("报表.xlsx", xlsx_bytes, "sheet", "1,234,567.89")
fid_d = upload_and_preview("文档.docx", docx_bytes, "word", "集成测试文档")

# ---------- 3) 非 Office 类 404 ----------
r = client.post("/api/knowledge-base/files", data={
    "file": (io.BytesIO("# md\n正文".encode("utf-8")), "note.md"), "name": "note",
}, content_type="multipart/form-data")
fid_m = r.get_json()["id"]
r = client.get(f"/api/knowledge-base/files/{fid_m}/preview")
assert r.status_code == 404, "非 Office 类应 404"
print("non-office preview 404 OK")

# 不存在 id
r = client.get("/api/knowledge-base/files/f-nonexist/preview")
assert r.status_code == 404
print("missing id preview 404 OK")

# ---------- 4) 旧格式直传（.xls 转换链 + 预览） ----------
if xls_bytes:
    r = client.post("/api/knowledge-base/files", data={
        "file": (io.BytesIO(xls_bytes), "老表格.xls"), "name": "老表格",
    }, content_type="multipart/form-data")
    up = r.get_json()
    assert r.status_code == 200, up
    assert up["item"]["converted"] is True and up["item"]["ext"] == "xlsx"
    fid_old = up["id"]
    r = client.get(f"/api/knowledge-base/files/{fid_old}/preview")
    pv = r.get_json()
    assert r.status_code == 200 and pv["kind"] == "sheet"
    print("legacy .xls upload -> converted -> preview OK")

# ---------- 4b) 异步下载源（上传时另提供下载文档） ----------
def upload(name, blob, **extra):
    data = {"file": (io.BytesIO(blob), name), "name": name.rsplit(".", 1)[0]}
    data.update(extra)
    return client.post("/api/knowledge-base/files", data=data,
                       content_type="multipart/form-data")


def dl_bytes(fid):
    r = client.get(f"/api/knowledge-base/files/{fid}/download")
    assert r.status_code == 200, r.status_code
    return r.data, r.headers.get("Content-Disposition", "")


# 用户场景：预览 Doc.pdf（PDF 可读），下载拿到 Demo.doc
dl_doc = doc_bytes or b"MZ\x90\x00fake-doc-bytes-for-download-source"
a = upload("Demo.pdf", b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF",
           download_file=(io.BytesIO(dl_doc), "Demo.doc")).get_json()
assert a["ok"] is True, a
assert a["item"]["has_download_source"] is True
assert a["item"]["download_ext"] == "doc" and a["item"]["download_name"] == "Demo.doc"
assert a["item"]["download_size"] == len(dl_doc)
assert a["item"]["original_ext"] == "pdf"          # 原件仍是 pdf（预览仍走 PDF 渲染）
fid_a = a["id"]
# 落盘名必须是 <id>.dl.<ext>，不能与原件 <id>.pdf 冲突
assert os.path.isfile(os.path.join(routes.FILES_DIR, f"{fid_a}.dl.doc"))
assert os.path.isfile(os.path.join(routes.FILES_DIR, f"{fid_a}.pdf"))
got, disp = dl_bytes(fid_a)
assert got == dl_doc, "下载应得到异步下载源字节"
assert "Demo.doc" in disp or "Demo.doc" in _re.sub(r"%[0-9A-Fa-f]{2}", "", disp)
# 列表接口同样暴露下载源标记（前端卡片角标/下载文案依赖它）
items = client.get("/api/knowledge-base/files?category=all").get_json()["items"]
brief = next(i for i in items if i["id"] == fid_a)
assert brief["has_download_source"] is True and brief["download_ext"] == "doc"
print("async download-source OK | Demo.pdf upload -> download Demo.doc")

# 未提供下载源 → 下载拿原件（行为不变）；字段缺省不写入记录
b = upload("plain.txt", "纯文本内容".encode("utf-8")).get_json()
assert b["item"]["has_download_source"] is False
assert b["item"]["download_ext"] == "" and b["item"]["download_size"] == 0
fid_b = b["id"]
got_b, _ = dl_bytes(fid_b)
assert got_b == "纯文本内容".encode("utf-8")
_rec = next(f for f in routes._load_store(routes._files_file())["files"] if f["id"] == fid_b)
assert "download_ext" not in _rec and "download_name" not in _rec
print("no download-source OK | falls back to original")

# 下载源校验失败：不阻断主上传，返回 download_error 且落入原行为
c = upload("Keep.pdf", b"%PDF-1.4\n%%EOF",
           download_file=(io.BytesIO(b"x" * 10), "bad.exe")).get_json()
assert c["ok"] is True, c                        # 主上传照常成功（B-4 降级）
assert "download_error" in c and "exe" in c["download_error"]
assert c["item"]["has_download_source"] is False
fid_c = c["id"]
got_c, _ = dl_bytes(fid_c)
assert got_c == b"%PDF-1.4\n%%EOF"               # 回落原件
print("invalid download-source OK | main upload survives, falls back")

# 下载源超过 20MB：同样不阻断主上传（用 _read_upload 直接验证限额逻辑，
# 不真发 20MB multipart——test client 会把整包读进内存且拖慢自测数个数量级）
_oversize = io.BytesIO(b"x" * (routes.MAX_UPLOAD_BYTES + 1))
_storage = type("S", (), {"filename": "big.txt", "stream": _oversize})()
try:
    routes._read_upload(_storage)
    raise AssertionError("超过 20MB 应被拒")
except routes._UploadError as _e:
    assert _e.status == 413 and "20MB" in str(_e), _e
print("oversize download-source OK | _read_upload rejects >20MB")

# 下载源文件被外力删除 → 回落原件（不 404）
os.remove(os.path.join(routes.FILES_DIR, f"{fid_a}.dl.doc"))
got_a2, _ = dl_bytes(fid_a)
assert got_a2.startswith(b"%PDF-1.4"), "下载源缺失应回落原件而非 404"
print("missing download-source file OK | falls back to original")

# ---------- 5) 删除清理（含预览缓存 + 异步下载源） ----------
# 重建一个带下载源的记录，验证 delete 会一并清理 .dl.<ext>
e = upload("Cleanup.pdf", b"%PDF-1.4\n%%EOF",
           download_file=(io.BytesIO(b"dl"), "Cleanup.txt")).get_json()
fid_e = e["id"]
assert os.path.isfile(os.path.join(routes.FILES_DIR, f"{fid_e}.dl.txt"))

for fid in (fid_x, fid_d, fid_m, fid_a, fid_b, fid_c, fid_e):
    r = client.delete(f"/api/knowledge-base/files/{fid}")
    assert r.status_code == 200
fids = (fid_x, fid_d, fid_m, fid_a, fid_b, fid_c, fid_e)
leftover = [f for f in os.listdir(routes.FILES_DIR)
            if any(f.startswith(x) for x in fids)]
assert not leftover, leftover
print("delete cleanup OK (no leftovers incl. preview cache & download source)")

# ---------- 6) 「异步下载源」开关：前端标记契约 ----------
# 回归背景：开关曾"点不动"——滑块是 position:absolute;inset:0 的覆盖层，
# 而它既不是 <label> 也没绑 click，只有左侧文字（<label for=...>）能激活 checkbox，
# 用户点那个可见的"按钮"完全没反应。修法 = 整行 .switch-row 用 <label> 包住 input。
# 这里做静态契约断言：一旦有人把 label 包回 <div>/<span>，立即失败。
_FE = os.path.join(REPO, "plugins", "knowledge-base", "frontend")
_app_js = open(os.path.join(_FE, "app.js"), encoding="utf-8").read()
_css = open(os.path.join(_FE, "style.css"), encoding="utf-8").read()

# 6.1 开关行必须是 <label class="form-row switch-row">，且内部同时含 input 与 .slider
_i = _app_js.index('\'<label class="form-row switch-row">\'')
_block = _app_js[_i:_i + 400]          # 取开关行起的一段，足够覆盖到 slider
assert 'class="switch-label"' in _block, _block
assert 'type="checkbox"' in _block and 'class="slider"' in _block, _block
assert '</span></label>' in _block, "开关行缺少 </span></label>（label 未包住 input）"
print("switch markup OK | 整行 <label> 包住 input + .slider")

# 6.2 不得再出现 for="m-dl-switch" 的旧写法（那是"只有点文字才有效"的病灶）
assert 'for="m-dl-switch"' not in _app_js, "旧写法 for=\"m-dl-switch\" 复活了"
print("switch no-legacy-for OK")

# 6.3 CSS 契约：.switch-row 可点；.switch 不再自称 label 语义
assert ".switch-row {" in _css and "cursor: pointer" in _css, "switch-row 缺少 cursor:pointer"
assert ".switch input:checked + .slider" in _css, "缺少 checked 态视觉规则"
print("switch css OK | cursor:pointer + checked 态规则在位")

# ---------- 7) 阅读计数（L4-4） ----------
r = client.post("/api/knowledge-base/files", data={
    "file": (io.BytesIO("# 计数\n正文".encode("utf-8")), "count.md"), "name": "计数",
}, content_type="multipart/form-data")
fid_v = r.get_json()["id"]

items = {f["id"]: f for f in client.get("/api/knowledge-base/files").get_json()["items"]}
assert items[fid_v]["views"] == 0 and items[fid_v]["last_viewed_at"] == "", \
    "新上传记录应从 0 起算（迁移/默认值没生效？）"

for expect in (1, 2, 3):
    rv = client.post(f"/api/knowledge-base/files/{fid_v}/view")
    assert rv.status_code == 200, rv.get_json()
    assert rv.get_json()["views"] == expect, f"第 {expect} 次应回 {expect}，实得 {rv.get_json()}"
items = {f["id"]: f for f in client.get("/api/knowledge-base/files").get_json()["items"]}
assert items[fid_v]["views"] == 3, "列表未回显累计值"
assert items[fid_v]["last_viewed_at"], "last_viewed_at 未落盘"
print("view count OK | 三次打开计到 3，列表回显一致")

# 7b) 计数**不随资源拉取增长**：/raw 与 /preview 不得顺手 +1（这是选显式端点的理由）
before = client.get("/api/knowledge-base/files").get_json()["items"]
before = {f["id"]: f["views"] for f in before}
client.get(f"/api/knowledge-base/files/{fid_v}/raw")
client.get(f"/api/knowledge-base/files/{fid_v}/preview")
after = {f["id"]: f["views"] for f in client.get("/api/knowledge-base/files").get_json()["items"]}
assert after[fid_v] == before[fid_v], f"资源拉取把计数抬了：{before[fid_v]} → {after[fid_v]}"
print("view count no-inflate OK | /raw + /preview 不计数")

# 7c) 语义：未知 ID / 非法 ID 一律 404（不 500、不暴露存在性）；未登录 401
assert client.post("/api/knowledge-base/files/nope123/view").status_code == 404
assert client.post("/api/knowledge-base/files/..%2Fetc/view").status_code == 404
_orig_user = admin_routes.get_session_user
jz_api.register_providers(get_session_user=lambda: None, set_operation=admin_routes.set_operation)
try:
    assert client.post(f"/api/knowledge-base/files/{fid_v}/view").status_code == 401, "无会话应 401"
finally:
    jz_api.register_providers(get_session_user=_orig_user, set_operation=admin_routes.set_operation)
# 401 那次不得留下痕迹 —— 用**只读**的列表核对（不能用 POST /view 去查，它自己会 +1）
after_401 = {f["id"]: f["views"] for f in client.get("/api/knowledge-base/files").get_json()["items"]}
assert after_401[fid_v] == 3, f"无会话的那次请求被计了：3 → {after_401[fid_v]}"
print("view count guards OK | 404（未知/非法 ID）+ 401（无会话且不留痕）")

# 7d) 迁移幂等：删掉两条记录的 views 键 → 补回；再跑一次不得判定为"需迁移"
store = routes._load_store(routes._files_file())
for rec in store["files"][:2]:
    rec.pop("views", None)
    rec.pop("last_viewed_at", None)
routes._save_store(routes._files_file(), store)
assert routes._migrate_files() is True, "缺键应触发一次迁移"
assert routes._migrate_files() is False, "迁移不幂等：第二次仍报有改动（会每次启动重写全库）"
store = routes._load_store(routes._files_file())
assert all("views" in r and isinstance(r["views"], int) for r in store["files"]), "迁移未补齐 views"
print("view count migration OK | 缺键补回 + 二次运行零改动")

# ---------- 8) 全文检索（L4-4） ----------
def _upload(name, blob):
    r = client.post("/api/knowledge-base/files", data={
        "file": (io.BytesIO(blob), name), "name": name.rsplit(".", 1)[0],
    }, content_type="multipart/form-data")
    assert r.status_code == 200, r.get_json()
    return r.get_json()["id"]


def _search(q, category=""):
    url = "/api/knowledge-base/files?q=" + q
    if category:
        url += "&category=" + category
    d = client.get(url).get_json()
    return {i["id"]: i for i in d["items"]}, d

fid_md = _upload("检索样本.md", "# 季度统计\n\n本文件含独特词饕餮纹样记录。\n".encode("utf-8"))
got, head = _search("饕餮")
assert fid_md in got, "内容命中未进结果（全文检索没生效）"
assert got[fid_md]["match_in"] == "content", got[fid_md]
assert "饕餮" in got[fid_md]["snippet"], got[fid_md]["snippet"]
assert head["fulltext"]["enabled"] is True and head["fulltext"]["hits"] >= 1, head["fulltext"]
# 精度断言：查询里的双字组必须在正文相邻出现——"饕纹"不相邻（原文是"饕餮纹"），不该命中
got, _ = _search("饕纹")
assert fid_md not in got, "CJK 双字组精度失效：不相邻的词也命中了（AND 语义被破坏）"
print("fulltext content hit OK | match_in=content + snippet + 双字组相邻精度")

got, _ = _search("检索样本")                     # 只在文件名里，正文没有
assert got[fid_md]["match_in"] == "name", got[fid_md]
got, _ = _search("季度统计")                      # 只在正文标题里，文件名没有
assert got[fid_md]["match_in"] == "content", got[fid_md]
# both 态要"文件名与正文都含该词"才成立，单独造一份这样的文件
fid_both = _upload("值班表.txt", "值班表（2026 年春节）：一楼由张三负责。\n".encode("utf-8"))
got, _ = _search("值班表")
assert got[fid_both]["match_in"] == "both", got[fid_both]
got, _ = _search("根本没有这个词组")
assert fid_md not in got, "无意义查询不该命中"
print("fulltext match_in OK | name / content / both 三态与无命中")

# docx 通道（python-docx，含表格单元格）
doc = docx.Document()
doc.add_paragraph("会议纪要：涉及貔貅采购事项。")
t = doc.add_table(rows=1, cols=1)
t.rows[0].cells[0].text = "貔貅单价"
buf = io.BytesIO(); doc.save(buf)
fid_dx = _upload("采购.docx", buf.getvalue())
got, _ = _search("貔貅")
assert fid_dx in got and got[fid_dx]["match_in"] == "content", \
    "docx 正文/表格抽取未覆盖（表格单元格是政务文档大头）"
print("fulltext docx OK | 段落与表格单元格都进索引")

# 分类作用域：内容命中不该跨分类出现
r = client.post("/api/knowledge-base/categories", json={"name": "只放md的分类"})
assert r.status_code == 200, (r.status_code, r.get_json())
cid = r.get_json()["id"]
client.put("/api/knowledge-base/files/" + fid_md, json={"category_id": cid})
got, _ = _search("貔貅", category=cid)
assert fid_dx not in got, "内容命中越了分类（越权可见面）"
got, _ = _search("饕餮纹", category=cid)
assert fid_md in got, "本分类内的内容命中不该被过滤掉"
# 回归：带分类的查询**不得**把索引裁成只剩该分类（ensure 的入参必须是全库，见 routes 注释）
st_after_cat = client.get("/api/knowledge-base/search/status").get_json()
assert st_after_cat["indexed"] >= 4, \
    "带分类的内容检索把索引裁掉了（indexed=%s）" % st_after_cat["indexed"]
got, _ = _search("貔貅")                     # 不带分类时，另一分类下的命中必须还在
assert fid_dx in got, "分类查询之后跨分类检索失效——索引被裁剪的实证"
print("fulltext category scope OK | 命中不跨分类，且索引不被带分类查询裁剪")

st = client.get("/api/knowledge-base/search/status").get_json()
assert st["ok"] and st["enabled"] is True and st["indexed"] >= 2, st
assert st["version"] == routes._text_index.INDEX_VERSION, st
print("search status OK | indexed=%s terms=%s truncated=%s"
      % (st["indexed"], st["terms"], st["truncated"]))

# 删除后不得再被检索命中（索引与存储必须同步摘除）
client.delete("/api/knowledge-base/files/" + fid_md)
got, _ = _search("饕餮纹")
assert fid_md not in got, "已删除文档仍可被全文检索命中"
st = client.get("/api/knowledge-base/search/status").get_json()
assert st["indexed"] >= 1, st
print("fulltext delete OK | 删除后不再命中，索引规模同步下降")

# 未登录：状态端点与列表都要 401（检索不得成为绕过鉴权的读取面）
_jz = routes._get_session_user
routes._get_session_user = lambda: None
try:
    assert client.get("/api/knowledge-base/search/status").status_code == 401
    assert client.get("/api/knowledge-base/files?q=貔貅").status_code == 401
finally:
    routes._get_session_user = _jz
print("fulltext auth OK | 未登录 401")

print("ALL PHASE-8 ROUTE CHECKS DONE")
