# -*- coding: utf-8 -*-
"""xlsx_render 自测：构造多场景 xlsx → 渲染 → 断言 + 落盘中间产物供浏览器目检。"""
import io
import json
import os
import sys
import datetime as dt

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "backend"))
import xlsx_render  # noqa: E402
import openpyxl  # noqa: E402
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "out")
os.makedirs(OUT, exist_ok=True)

wb = openpyxl.Workbook()

# ---------- Sheet1：综合样式 ----------
ws = wb.active
ws.title = "综合样式"

# 标题行（合并 + 主题色填充 + 白字）
ws.merge_cells("A1:G1")
c = ws["A1"]
c.value = "2026年三季度销售汇总表"
c.font = Font(name="微软雅黑", size=16, bold=True, color="FFFFFF")
c.fill = PatternFill("solid", fgColor="4472C4")
c.alignment = Alignment(horizontal="center", vertical="center")
ws.row_dimensions[1].height = 32

# 表头
headers = ["序号", "产品", "数量", "单价", "金额", "完成率", "日期"]
for i, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=i, value=h)
    c.font = Font(bold=True, size=11)
    c.fill = PatternFill("solid", fgColor="D9E1F2")
    c.border = Border(top=Side("thin"), bottom=Side("double"),
                      left=Side("thin"), right=Side("thin"))
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[2].height = 24

# 数据行
data = [
    (1, "铅酸电池", 1200, 89.5, 107400, 0.856, dt.date(2026, 7, 3)),
    (2, "锂电池组", 350, 1250.0, 437500, 1.12, dt.date(2026, 7, 18)),
    (3, "充电器", 5800, 39.9, 231420, 0.4325, dt.date(2026, 8, 2)),
    (4, "逆变器", -12, 4200.0, -50400, 0.75, dt.date(2026, 8, 21)),
    (5, "线材", 9999999, 2.5, 24999997.5, 0.999, dt.datetime(2026, 9, 1, 14, 30, 5)),
]
thin = Side("thin", color="808080")
for r, row in enumerate(data, 3):
    for i, v in enumerate(row, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.border = Border(bottom=thin, right=Side("hair", color="B0B0B0"))
        c.alignment = Alignment(vertical="center")
    ws.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="center")
    ws.cell(row=r, column=3).number_format = "#,##0"
    ws.cell(row=r, column=4).number_format = "#,##0.00"
    ws.cell(row=r, column=5).number_format = "¥#,##0.00;[Red]-¥#,##0.00"
    ws.cell(row=r, column=6).number_format = "0.00%"
    ws.cell(row=r, column=7).number_format = 'yyyy"年"m"月"d"日"'

# 合计行（跨列 + 主题色 + tint）
ws.merge_cells("A8:E8")
c = ws["A8"]
c.value = "合计（数字格式 0.0 测试：1234.5678）"
c.number_format = "0.0"
c.font = Font(bold=True)
c.alignment = Alignment(horizontal="right", vertical="center")
for i in range(1, 6):
    ws.cell(row=8, column=i).border = Border(top=Side("medium"), bottom=Side("medium"))
c2 = ws.cell(row=8, column=6, value=0.5678)
c2.number_format = "0.0%"
c2.fill = PatternFill("solid", fgColor="FFC000")
c2.alignment = Alignment(horizontal="center")

# 隐藏行列 + 列宽
ws.column_dimensions["A"].width = 6
ws.column_dimensions["B"].width = 14
ws.column_dimensions["D"].width = 10
ws.column_dimensions["G"].width = 14
ws.row_dimensions[5].hidden = True
# 主题色 + tint（fgColor 用 theme/tint 而非 rgb）
from openpyxl.styles.colors import Color  # noqa: E402
c = ws.cell(row=10, column=4, value="主题色 accent1+tint0.6")
c.fill = PatternFill("solid", fgColor=Color(theme=4, tint=0.5999938962981048))
c.font = Font(color=Color(theme=1, tint=0.499984740745262))
c.alignment = Alignment(horizontal="center", vertical="center")
# 可见货币列 + [Red] 负数
ws["D11"] = 107400
ws["D11"].number_format = "¥#,##0.00"
ws["D12"] = -50400
ws["D12"].number_format = "¥#,##0.00;[Red]-¥#,##0.00"
ws["D12"].font = Font(bold=True)

# 隐藏列（有内容）
ws["H2"] = "本列隐藏"
ws["H3"] = 999
ws.column_dimensions["H"].hidden = True

# 时间格式
ws["C10"] = dt.time(9, 5, 3)
ws["C10"].number_format = "h:mm AM/PM"
ws["C11"] = dt.time(15, 42, 10)
ws["C11"].number_format = "hh:mm:ss"
ws["C12"] = dt.timedelta(hours=31, minutes=25)
ws["C12"].number_format = "[h]:mm"
ws["C13"] = 0.5
ws["C13"].number_format = "# ?/?"   # 分数格式（近似处理）

# ---------- Sheet2：中文宽表（模拟台账） ----------
ws2 = wb.create_sheet("台账")
ws2["A1"] = "设备台账（关闭网格线测试）"
ws2["A1"].font = Font(size=14, bold=True, underline="single", color="C00000")
long_text = "这是一段很长很长的备注文本用于测试自动换行是否按照 Excel 的 wrap_text 行为正确折行显示而不溢出。"
ws2["A2"] = long_text
ws2["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws2.merge_cells("A2:D2")
ws2.row_dimensions[2].height = 40
ws2.sheet_view.showGridLines = False
for i, w in enumerate((20, 12, 12, 30, 10), 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2["A4"] = 1234567.891
ws2["A4"].number_format = "#,##0.00"
ws2["B4"] = -0.1234
ws2["B4"].number_format = "0.00%;[Red]0.00%"
ws2["C4"] = 1234567890123
ws2["C4"].number_format = "General"
ws2["D4"] = "文本型数字 007"
ws2["E4"] = 0.000000123
ws2["E4"].number_format = "0.00E+00"

# ---------- Sheet3：空表 ----------
wb.create_sheet("空表")

# ---------- Sheet4：超宽表（触发 MAX_COLS=120 截断提示） ----------
ws4 = wb.create_sheet("超宽表")
ws4.cell(row=1, column=1, value="截断提示测试")
for cidx in range(1, 130):
    ws4.cell(row=2, column=cidx, value="C%d" % cidx)

xlsx_path = os.path.join(OUT, "sample.xlsx")
wb.save(xlsx_path)

summary = xlsx_render.render_xlsx(xlsx_path, os.path.join(OUT, "sample.json"))
with open(os.path.join(OUT, "sample.json"), encoding="utf-8") as f:
    payload = json.load(f)

html_path = os.path.join(OUT, "sample.html")
with open(html_path, "w", encoding="utf-8") as f:
    f.write("""<!DOCTYPE html><html><head><meta charset="utf-8"><title>预览目检</title>
<style>body{font-family:sans-serif;background:#f0f2f5;margin:0;padding:16px}
.sheet-tabs{display:flex;gap:8px;margin-bottom:12px}
.sheet-tab{padding:4px 14px;border:1px solid #ccc;border-radius:6px;background:#fff;cursor:pointer;font-size:13px}
.sheet-tab.active{background:#2f6fbf;border-color:#2f6fbf;color:#fff}
.sheet-table-wrap{overflow:auto;max-height:80vh;background:#fff;padding:12px;border-radius:8px}
.kb-xlsx-table td{border:1px solid #d9d9d9;padding:0 3px}
.kb-nogrid td{border-color:transparent}
.kb-xlsx-more td{text-align:center;background:#fffbe6;color:#8a6d3b}
</style></head><body><div id="app"></div>
<script>
var data = %s;
var tabs = document.createElement('div'); tabs.className='sheet-tabs';
var wrap = document.createElement('div'); wrap.className='sheet-table-wrap';
function show(i){ wrap.innerHTML = data.sheets[i].html;
  Array.from(tabs.children).forEach(function(b,j){ b.className='sheet-tab'+(j===i?' active':''); }); }
data.sheets.forEach(function(s,i){ var b=document.createElement('button');
  b.className='sheet-tab'; b.textContent=s.name; b.onclick=function(){show(i);}; tabs.appendChild(b); });
document.getElementById('app').appendChild(tabs); document.getElementById('app').appendChild(wrap);
show(0);
</script></body></html>""" % json.dumps(payload, ensure_ascii=False))

print("summary:", summary)
print("sheets:", [(s["name"], s["rows_total"], s["cols_total"], s["truncated"]) for s in payload["sheets"]])
