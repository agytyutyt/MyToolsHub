# -*- coding: utf-8 -*-
"""重建阶段8目检页（完全复刻 reader.js 注入逻辑，含 style 摘取挂回）。"""
import io
import json
import os
import sys

BACKEND = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, BACKEND)
sys.path.insert(0, os.path.join(BACKEND, "vendor"))
import office_render  # noqa: E402
import docx  # noqa: E402

with open(os.path.join(BACKEND, "out", "sample.xlsx"), "rb") as f:
    sheet = office_render.render_sheet(f.read())

d = docx.Document()
d.add_heading("季度工作报告", level=0)
d.add_paragraph("本报告由渲染引擎自动转换，用于验证 Word 预览效果。")
d.add_heading("一、销售情况", level=1)
p = d.add_paragraph()
p.add_run("重点区域").bold = True
p.add_run("实现销售额 1,234.5 万元，同比增长 12%；其他区域持平。")
d.add_heading("二、明细", level=1)
t = d.add_table(rows=3, cols=3)
t.style = "Table Grid"
for j, h in enumerate(("产品", "数量", "金额")):
    t.rows[0].cells[j].text = h
for i, row in enumerate((("铅酸电池", "1,200", "107.4"), ("锂电池组", "350", "437.5")), 1):
    for j, v in enumerate(row):
        t.rows[i].cells[j].text = v
d.add_paragraph("三、备注", style="List Number")
d.add_paragraph("数据口径以财务为准。", style="List Bullet")
buf = io.BytesIO()
d.save(buf)
word = office_render.render_word(buf.getvalue())

NL = chr(10)
BS_BS = chr(92)  # 反斜杠
JOIN_NL = 'cssParts.join("' + BS_BS + 'n")'

page = """<!DOCTYPE html><html><head><meta charset="utf-8"><title>阶段8目检</title>
<link rel="stylesheet" href="../../frontend/style.css?v=22">
<style>body{background:#eef1f5;font-family:sans-serif;margin:0;padding:12px}
h4{margin:14px 0 6px}
#reader-container,#reader-container2{background:#fff;border-radius:8px;padding:12px;max-width:1100px}
</style></head><body>
<h4>Word（dhr 引擎 → style 摘取 + DOMPurify → 注入）</h4>
<div id="reader-container"><div id="w1"></div></div>
<h4>Excel（xhr 引擎 → style 摘取 + DOMPurify → 注入 + 页签接线）</h4>
<div id="reader-container2"><div id="w2"></div></div>
<script src="../../frontend/vendor/purify/purify.min.js?v=1"></script>
<script>
var DATA = %(payload)s;
function inject(el, data, kind){
  var t0 = performance.now();
  var cssParts = [];
  var body = data.html.replace(/<style[^>]*>([%(s)s%(S)s]*?)<%(sl)s/style>/gi,
    function(m0, css){ cssParts.push(css); return ""; });
  var clean = window.DOMPurify.sanitize(body, { USE_PROFILES: { html: true } });
  var ms = (performance.now()-t0).toFixed(1);
  var box = document.createElement("div");
  box.className = "kb-office" + (kind==="sheet" ? " kb-office-sheet" : " kb-office-word");
  box.innerHTML = clean;
  if (cssParts.length) {
    var st = document.createElement("style");
    st.textContent = %(join)s;
    box.insertBefore(st, box.firstChild);
  }
  el.appendChild(box);
  var info = document.createElement("p");
  info.textContent = "sanitize " + ms + "ms | style块摘取: " + cssParts.length + " | 原始 " + data.html.length + "B";
  el.insertBefore(info, box);
  if (kind === "sheet") {
    box.addEventListener("click", function (ev) {
      var t = ev.target;
      if (!t || !t.className || String(t.className).indexOf("kbsheet-tab") < 0) return;
      var idx = t.getAttribute("data-target");
      var tabs = box.querySelectorAll(".kbsheet-tab");
      var sheets = box.querySelectorAll(".kbsheet-sheet");
      for (var i = 0; i < tabs.length; i++) tabs[i].setAttribute("aria-selected", String(i === Number(idx)));
      for (var j = 0; j < sheets.length; j++) sheets[j].hidden = String(j) !== idx;
    });
  }
  return cssParts.length > 0;
}
var s1 = inject(document.getElementById("w1"), DATA.word, "word");
var s2 = inject(document.getElementById("w2"), DATA.sheet, "sheet");
document.title = "STYLE_OK=" + (s1 && s2);
</script></body></html>""" % {
    "payload": json.dumps({"sheet": sheet, "word": word}, ensure_ascii=False),
    "s": BS_BS + "s", "S": BS_BS + "S", "sl": BS_BS, "join": JOIN_NL,
}

out = os.path.join(BACKEND, "out", "phase8.html")
with open(out, "w", encoding="utf-8") as f:
    f.write(page)
print("rebuilt:", out, "| sheet %dB word %dB" % (len(sheet["html"]), len(word["html"])))
